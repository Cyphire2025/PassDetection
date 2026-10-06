"""Exercise actual image math, PDF rendering and OCR inside the final runtime."""
import argparse
import importlib.metadata
import importlib.util
import json
import re
import subprocess
import sys
from pathlib import Path

import cv2
import matplotlib
import numpy as np
import passporteye
import pypdfium2
import pytesseract
from PIL import Image, ImageDraw, ImageFont


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--require-no-video", action="store_true")
    parser.add_argument("--require-patched-parsers", action="store_true")
    args = parser.parse_args()
    sys.path.insert(0, str(Path.cwd()))
    font = ImageFont.truetype(str(Path(matplotlib.get_data_path()) / "fonts/ttf/DejaVuSans.ttf"), 48)
    image = Image.new("RGB", (900, 160), "white")
    ImageDraw.Draw(image).text((20, 30), "PASSPORT 12345678", font=font, fill="black")
    text = pytesseract.image_to_string(image, config="--psm 7", timeout=20)
    if "PASSPORT" not in text or "12345678" not in text:
        raise RuntimeError("Synthetic OCR result differs from the fixture")
    array = np.asarray(image.convert("L"))
    threshold = cv2.threshold(cv2.GaussianBlur(array, (3, 3), 0), 0, 255, cv2.THRESH_BINARY_INV | cv2.THRESH_OTSU)[1]
    contours, _ = cv2.findContours(threshold, cv2.RETR_EXTERNAL, cv2.CHAIN_APPROX_SIMPLE)
    if len(contours) < 10:
        raise RuntimeError("Synthetic image processing lost text contours")
    from app.infrastructure.ocr.mrz_image_normalizer import MRZImageNormalizer
    normalized = MRZImageNormalizer().normalize(image)
    if normalized is None:
        raise RuntimeError("MRZ normalization returned no image")
    if importlib.util.find_spec("setuptools") or importlib.util.find_spec("wheel"):
        raise RuntimeError("Build-only tooling survived in the runtime environment")
    # Create/render a synthetic PDF through the same supported image libraries.
    import io
    pdf = io.BytesIO()
    image.save(pdf, format="PDF")
    document = pypdfium2.PdfDocument(pdf.getvalue())
    page = document[0]
    bitmap = page.render(scale=1)
    if bitmap.width <= 0 or bitmap.height <= 0:
        raise RuntimeError("Synthetic PDF rendering failed")
    bitmap.close()
    page.close()
    document.close()
    video = hasattr(cv2, "VideoCapture")
    ffmpeg = bool(re.search(r"FFMPEG:\s+YES", cv2.getBuildInformation()))
    if args.require_no_video and (video or ffmpeg):
        raise RuntimeError("Unneeded video APIs or FFmpeg survived the restricted image build")
    parser_checks = {}
    if args.require_patched_parsers:
        security_packages = json.loads(subprocess.check_output([
            sys.executable, "/usr/local/share/passdetection/verify_runtime_security_packages.py",
        ], text=True))
        import io
        import pyexpat
        import xml.etree.ElementTree as ET

        import openpyxl
        for encoding in ("utf-16-le", "utf-16-be"):
            malformed = "<r>\ud800a</r>".encode(encoding, errors="surrogatepass")
            try:
                pyexpat.ParserCreate(encoding="UTF-16").Parse(malformed, True)
            except pyexpat.ExpatError:
                pass
            else:
                raise RuntimeError("Malformed UTF-16 surrogate accepted by the actual Python parser")
            try:
                ET.fromstring(malformed)
            except ET.ParseError:
                pass
            else:
                raise RuntimeError("ElementTree still reaches an unpatched Expat copy")
        workbook = openpyxl.Workbook()
        workbook.active.append(["Synthetic passport", "L898902C3", 42])
        stream = io.BytesIO()
        workbook.save(stream)
        workbook.close()
        stream.seek(0)
        restored = openpyxl.load_workbook(stream, read_only=True)
        if next(restored.active.values) != ("Synthetic passport", "L898902C3", 42):
            raise RuntimeError("XLSX round trip changed after the XML parser repair")
        restored.close()
        import tarfile
        import tempfile
        # Upstream CVE-2026-82049 fixture: relocating a hard link to a relative
        # symlink must resolve the safe original file, not relocate the symlink.
        archive = io.BytesIO()
        with tarfile.open(fileobj=archive, mode="w") as writer:
            regular = tarfile.TarInfo("a/escape")
            regular.size = 5
            writer.addfile(regular, io.BytesIO(b"decoy"))
            symlink = tarfile.TarInfo("a/b/s")
            symlink.type, symlink.linkname = tarfile.SYMTYPE, "../escape"
            writer.addfile(symlink)
            hardlink = tarfile.TarInfo("s")
            hardlink.type, hardlink.linkname = tarfile.LNKTYPE, "a/b/s"
            writer.addfile(hardlink)
        for extraction_filter in ("data", "tar"):
            with tempfile.TemporaryDirectory() as directory:
                with tarfile.open(fileobj=io.BytesIO(archive.getvalue())) as reader:
                    reader.extractall(directory, filter=extraction_filter)
                result = Path(directory) / "s"
                if result.is_symlink() or result.read_bytes() != b"decoy":
                    raise RuntimeError("tarfile still relocates an unsafe hard-linked symlink")
        binary = subprocess.check_output(["sh", "-c", "command -v tesseract"], text=True).strip()
        linkage = subprocess.check_output(["ldd", binary], text=True)
        if any(name in linkage for name in ("libcurl", "libarchive", "libxml2")):
            raise RuntimeError("Unneeded network/archive/XML OCR parsers remain linked")
        languages = subprocess.check_output(["tesseract", "--list-langs"], text=True)
        if not {"eng", "osd"} <= set(languages.splitlines()):
            raise RuntimeError("Required English and orientation OCR data is missing")
        parser_checks = {**security_packages, "expat": pyexpat.EXPAT_VERSION, "utf16_negatives": 4,
                         "tarfile_hardlink_relocation_negatives": 2,
                         "xlsx_round_trip": True, "ocr_url_archive_dependencies_absent": True,
                         "ocr_languages": ["eng", "osd"], "tesseract": str(pytesseract.get_tesseract_version())}
    print(json.dumps({"result": "passed", "opencv": cv2.__version__, "numpy": np.__version__,
        "ocr_text": text.strip(), "contours": len(contours), "pdf_render": True,
        "passporteye_import": bool(passporteye), "runtime_build_tools_absent": True,
        "ffmpeg_enabled": ffmpeg,
        "video_api_present": video, "parser_checks": parser_checks,
        "python_dependencies": {name: importlib.metadata.version(name) for name in ("pillow", "pypdfium2", "passporteye")}}))


if __name__ == "__main__":
    main()
