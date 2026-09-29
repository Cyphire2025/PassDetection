"""Bounded MCP traffic and evidence for the fixed synthetic capacity lane.

Tokens, artifact handles, filenames and response bodies never enter samples.
This exercises the actual authenticated HTTP boundary; fixture sign-in is not
evidence of a browser/OS-vault or real-Codex authentication journey.
"""

from __future__ import annotations

import asyncio
import hashlib
import io
import re
import time
from collections import Counter
from datetime import UTC, datetime

import httpx
from mcp_capacity_profile import EXPORT_SCHEDULING, SOURCE_ROWS
from openpyxl import load_workbook

MCP_BUDGETS_MS = {
    "mcp_groups": (1500, 3000),
    "mcp_roster": (2000, 4000),
    "mcp_export": (15000, 30000),
}


def headers(actor: dict) -> dict[str, str]:
    return {
        "Authorization": f"Bearer {actor['token']}",
        "Accept": "application/json, text/event-stream",
        "MCP-Protocol-Version": "2025-11-25",
    }


async def call(client: httpx.AsyncClient, actor: dict, name: str, arguments: dict) -> dict:
    response = await client.post(
        "/mcp", headers=headers(actor),
        json={"jsonrpc": "2.0", "id": 1, "method": "tools/call",
              "params": {"name": name, "arguments": arguments}},
    )
    response.raise_for_status()
    envelope = response.json()
    result = envelope.get("result", {})
    body = result.get("structuredContent", {})
    if (envelope.get("id") != 1 or "error" in envelope or result.get("isError")
            or not isinstance(body, dict) or "error" in body
            or not body.get("audit_id") or not body.get("observed_at")
            or body.get("completeness") not in {"complete", "partial"}):
        raise ValueError("MCP operation failed its bounded response contract")
    return body


async def measured_read(client: httpx.AsyncClient, samples: list, actor: dict, iteration: int) -> None:
    roster = iteration % 2 == 1
    operation = "mcp_roster" if roster else "mcp_groups"
    started, at = time.perf_counter(), datetime.now(UTC).isoformat()
    valid, count = False, 0
    try:
        body = await call(client, actor, "list_group_passports" if roster else "list_groups", {
            "group_id": actor["group_id"], "page_size": 50,
            **({} if roster else {"agency_id": actor["agency_id"]}),
        })
        rows = body["items"]
        count = len(rows)
        if roster:
            valid = (count == min(50, actor["group_size"])
                     and body["group_total"] == actor["group_size"]
                     and body["group_id"] == actor["group_id"])
        else:
            valid = (count == 1 and rows[0]["id"] == actor["group_id"]
                     and rows[0]["counts"]["passport_submission_records"] == actor["group_size"])
    except (httpx.HTTPError, ValueError, KeyError, TypeError):
        pass
    samples.append({
        "operation": operation, "started_at": at, "finished_at": datetime.now(UTC).isoformat(),
        "status": 200 if valid else 0, "milliseconds": (time.perf_counter() - started) * 1000,
        "valid": valid, "bytes": 0, "rows": count, "cohort": actor["cohort"],
    })


async def paced_reads(client: httpx.AsyncClient, samples: list, actor: dict, seconds: int) -> None:
    start = time.monotonic()
    for iteration in range(seconds):
        await asyncio.sleep(max(0, start + iteration - time.monotonic()))
        await measured_read(client, samples, actor, iteration)


def verify_workbook(data: io.BytesIO, actor: dict) -> int:
    """Require every intended synthetic traveller, including duplicate passports."""
    workbook = load_workbook(data, read_only=True, data_only=False)
    try:
        if workbook.sheetnames != ["Passport Submissions"]:
            raise ValueError("Unexpected workbook sheets")
        sheet = workbook["Passport Submissions"]
        columns = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
        names = ("GIVEN NAME", "SURNAME", "Passport Number")
        if any(columns.count(name) != 1 for name in names):
            raise ValueError("Export omitted the canonical traveller columns")
        indexes = [columns.index(name) for name in names]
        observed = Counter()
        for row in sheet.iter_rows(min_row=5, values_only=True):
            if not any(value is not None for value in row):
                continue
            values = tuple(str(row[index] or "").strip().casefold() for index in indexes)
            observed[values] += 1
        tenant, count = actor["tenant"], actor["export_size"]
        expected = Counter((f"synthetic {index}", f"traveller {tenant}",
                            (f"P{tenant:02d}{index:05d}" if index >= 50 else f"DUP{tenant:02d}").casefold())
                           for index in range(count))
        if observed != expected:
            raise ValueError("Export rows differ from the exact intended selection")
        return sum(observed.values())
    finally:
        workbook.close()


async def verified_export(client: httpx.AsyncClient, samples: list, actor: dict, run_id: str, stage: str) -> dict:
    started, at = time.perf_counter(), datetime.now(UTC).isoformat()
    valid, size, workbook_rows = False, 0, 0
    try:
        if (type(actor["export_size"]) is not int or not 0 < actor["export_size"] <= SOURCE_ROWS
                or (actor["group_size"] > SOURCE_ROWS
                    and len(actor["export_submission_ids"]) != actor["export_size"])):
            raise ValueError("Export exceeds the minimum deployed source envelope")
        selection = {"agency_id": actor["agency_id"], "group_ids": [actor["group_id"]]}
        if actor["export_submission_ids"]:
            selection.update(selection="selected_passports", submission_ids=actor["export_submission_ids"])
        inspected = await call(client, actor, "inspect_excel_export", {"export": selection})
        if (inspected.get("passenger_count") != actor["export_size"]
                or inspected.get("pending_recipient_count") != 0):
            raise ValueError("Inspected export differs from the explicit synthetic selection")
        prepared = await call(client, actor, "prepare_excel_export", {
            "export": selection, "expected_revision": inspected["expected_revision"],
            "idempotency_key": f"capacity-{run_id}-{stage}-{actor['cohort']}",
        })
        artifact = prepared["artifact"]
        handle = artifact["artifact_id"]
        if not isinstance(handle, str) or not re.fullmatch(r"[a-zA-Z0-9_-]{16,256}", handle):
            raise ValueError("Unexpected protected artifact locator")
        expected = artifact["byte_size"]
        if type(expected) is not int or not 0 < expected <= 32 * 1024 * 1024:
            raise ValueError("Capacity export exceeds its declared memory bound")
        data = io.BytesIO()
        digest = hashlib.sha256()
        async with client.stream("GET", f"/mcp/artifacts/{handle}/content", headers=headers(actor)) as response:
            response.raise_for_status()
            async for chunk in response.aiter_bytes(65536):
                size += len(chunk)
                if size > expected:
                    raise ValueError("Export exceeded authorized size")
                digest.update(chunk)
                data.write(chunk)
        if size != expected or digest.hexdigest() != artifact["sha256"]:
            raise ValueError("Export checksum mismatch")
        data.seek(0)
        workbook_rows = verify_workbook(data, actor)
        delivered = await client.post(f"/mcp/artifacts/{handle}/delivery", headers=headers(actor),
                                      json={"byte_size": size, "sha256": digest.hexdigest()})
        delivered.raise_for_status()
        metadata = delivered.json()
        valid = bool(metadata.get("delivered_at")) and metadata.get("sha256") == digest.hexdigest()
    except (httpx.HTTPError, ValueError, KeyError, TypeError, OSError):
        pass
    sample = {
        "operation": "mcp_export", "started_at": at, "finished_at": datetime.now(UTC).isoformat(),
        "status": 200 if valid else 0, "milliseconds": (time.perf_counter() - started) * 1000,
        "valid": valid, "bytes": size, "rows": workbook_rows, "cohort": actor["cohort"],
    }
    samples.append(sample)
    return sample


async def serialized_exports(client, samples, actors, run_id, stage):
    """One export journey at a time; planned wait still counts in latency gates.

    Website and MCP read traffic remain concurrent. Busy/error responses fail the
    sample; this lane does not silently retry, mint new keys, or claim contention
    qualification. Application admission and cancellation have separate tests.
    """
    started, at = time.perf_counter(), datetime.now(UTC).isoformat()
    for order, actor in enumerate(actors):
        wait_ms = (time.perf_counter() - started) * 1000
        sample = await verified_export(client, samples, actor, run_id, stage)
        sample.update(
            export_scheduling=EXPORT_SCHEDULING, export_order=order,
            execution_started_at=sample["started_at"], started_at=at,
            serialization_wait_ms=wait_ms,
            milliseconds=sample["milliseconds"] + wait_ms,
        )


def gates(samples: dict, windows: dict, metrics: list[dict]) -> list[str]:
    failures = []
    for stage in ("expected", "recovery"):
        rows = samples[stage]
        for cohort in ("mcp_large_group", "mcp_small_group"):
            reads = [r for r in rows if r["cohort"] == cohort and r["operation"] in {"mcp_groups", "mcp_roster"}]
            if len(reads) / windows[stage]["elapsed_seconds"] < 0.9:
                failures.append(f"{stage}:{cohort}:mcp_throughput_below_0.9rps")
            for operation in ("mcp_groups", "mcp_roster", "mcp_export"):
                selected = [r for r in rows if r["cohort"] == cohort and r["operation"] == operation]
                if not selected or any(not r["valid"] for r in selected):
                    failures.append(f"{stage}:{cohort}:{operation}:missing_or_invalid")
                if operation == "mcp_export" and (
                    len(selected) != 1 or any(
                        r.get("rows") != SOURCE_ROWS
                        or r.get("export_scheduling") != EXPORT_SCHEDULING
                        for r in selected
                    )
                ):
                    failures.append(f"{stage}:{cohort}:mcp_export:wrong_profile_or_schedule")
                values = sorted(r["milliseconds"] for r in selected)
                if values:
                    import math
                    if any(values[math.ceil(len(values) * percentile) - 1] > budget
                           for percentile, budget in zip((0.95, 0.99), MCP_BUDGETS_MS[operation], strict=True)):
                        failures.append(f"{stage}:{cohort}:{operation}:latency_budget")
    events = [row.get("backend_memory_events") for row in metrics]
    if not events or any(not isinstance(row, dict) or not {"oom", "oom_kill"} <= row.keys() for row in events):
        failures.append("mcp:missing_oom_evidence")
    elif any(any(row[key] != events[0][key] for key in ("oom", "oom_kill")) for row in events[1:]):
        failures.append("mcp:oom_event_during_workload")
    return failures
