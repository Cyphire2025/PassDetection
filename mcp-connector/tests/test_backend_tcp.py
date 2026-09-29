"""Actual local TCP qualification against the application's MCP/OAuth adapters.

Requires the repository backend test runtime. No browser, OS vault, production
database, external provider, or Codex configuration participates in this test.
"""

import asyncio
import json
import socket
import sys
import time
import uuid
from contextlib import asynccontextmanager
from datetime import UTC, datetime
from pathlib import Path

import httpx2
import pytest
import uvicorn
from fastapi import FastAPI
from mcp import Client
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

from gc_mcp_connector.auth import Authorization
from gc_mcp_connector.config import CLIENT_ID, REDIRECT_URI, Config
from gc_mcp_connector.oauth import AuthorizationAttempt, OAuthClient
from gc_mcp_connector.proxy import RemoteProxy
from gc_mcp_connector.vault import Credential


class FixtureVault:
    """In-memory fixture only; the production CLI has no such fallback."""

    value: Credential | None = None

    def read(self):
        return self.value

    def write(self, credential):
        self.value = credential

    def delete(self):
        self.value = None


@asynccontextmanager
async def fixture_lock():
    yield


@pytest.fixture
async def backend_tcp(tmp_path, monkeypatch):
    backend = Path(__file__).resolve().parents[2] / "backend"
    sys.path.insert(0, str(backend))
    # Nested backend settings classes also load .env relative to cwd; isolate it.
    monkeypatch.chdir(tmp_path)
    for key, value in {
        "APP_ENV": "development",
        "APP_SECRET_KEY": "connector-tcp-fixture-secret-not-for-production",
        "POSTGRES_PASSWORD": "connector-fixture-database-password",
        "S3_ACCESS_KEY_ID": "connector-fixture-storage-access",
        "S3_SECRET_ACCESS_KEY": "connector-fixture-storage-secret",
    }.items():
        monkeypatch.setenv(key, value)

    from app.application.mcp.authorization import MCPAuthorizationService
    from app.core.config.mcp import MCPSettings
    from app.core.config.settings import Settings
    from app.infrastructure.database.mcp_models import MCPControlModel
    from app.infrastructure.database.model_registry import Base
    from app.infrastructure.database.models import UserModel, UserSecurityStateModel
    from app.infrastructure.database.session import get_db_session
    from app.presentation.api.v1.routes.mcp_artifacts import router as artifact_router
    from app.presentation.api.v1.routes.mcp_oauth import router as oauth_router
    from app.presentation.api.v1.routes.mcp_whatsapp_media import router as whatsapp_media_router
    from app.presentation.mcp.server import install_mcp

    sock = socket.socket(socket.AF_INET, socket.SOCK_STREAM)
    sock.bind(("127.0.0.1", 0))
    origin = f"http://127.0.0.1:{sock.getsockname()[1]}"
    config = Config(origin)
    settings = Settings(
        app_env="development",
        app_secret_key="connector-tcp-fixture-secret-not-for-production",
        app_debug=False,
        mcp=MCPSettings(enabled=True, public_origin=origin, frontend_origin=origin),
        _env_file=None,
    )
    engine = create_async_engine(f"sqlite+aiosqlite:///{tmp_path / 'connector-fixture.sqlite'}")
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with engine.begin() as connection:
        await connection.run_sync(Base.metadata.create_all)
    now, user_id = datetime.now(UTC), uuid.uuid4()
    attempt = AuthorizationAttempt.create()
    async with sessions() as session:
        session.add(
            UserModel(
                id=user_id,
                email="connector-tcp@example.test",
                hashed_password="fixture",
                full_name="Connector TCP fixture",
                role="super_admin",
                is_active=True,
            )
        )
        await session.flush()
        session.add_all(
            [
                UserSecurityStateModel(
                    user_id=user_id,
                    session_version=1,
                    credential_state="active",
                    mfa_secret_ciphertext="fixture",
                    mfa_enabled_at=now,
                ),
                MCPControlModel(id=1, enabled=True),
            ]
        )
        await session.flush()
        code = await MCPAuthorizationService(session, settings).authorize(
            user_id=user_id,
            security_version=1,
            mfa_at=now,
            client_id=CLIENT_ID,
            redirect_uri=REDIRECT_URI,
            resource=config.resource,
            challenge=attempt.challenge,
            scopes=["mcp:read", "mcp:upload", "mcp:export"],
            name="Isolated connector TCP fixture",
        )
        await session.commit()

    app = FastAPI()
    app.state.settings = settings
    app.include_router(oauth_router)
    app.include_router(artifact_router)
    app.include_router(whatsapp_media_router)
    install_mcp(app, settings)
    app.state.mcp_session_factory = sessions

    async def session_override():
        async with sessions() as session:
            yield session

    app.dependency_overrides[get_db_session] = session_override
    server = uvicorn.Server(
        uvicorn.Config(app, log_level="error", access_log=False, lifespan="on", ws="none")
    )
    task = asyncio.create_task(server.serve(sockets=[sock]))
    try:
        async with asyncio.timeout(10):
            while not server.started:
                if task.done():
                    await task
                await asyncio.sleep(0.01)
        yield config, sessions, attempt, code, user_id, app
    finally:
        server.should_exit = True
        await asyncio.wait_for(task, timeout=10)
        sock.close()
        await engine.dispose()


async def test_browser_code_exchange_refresh_sdk_proxy_and_revocation_over_real_tcp(backend_tcp):
    from app.infrastructure.database.mcp_models import MCPGrantModel, MCPTokenModel
    from app.infrastructure.database.models import AuditLogModel

    config, sessions, attempt, code, _user_id, _app = backend_tcp
    async with httpx2.AsyncClient(timeout=10, trust_env=False, follow_redirects=False) as http:
        oauth = OAuthClient(config, http)
        await oauth.discover()
        tokens = await oauth.exchange(code, attempt)
        vault = FixtureVault()
        authorization = Authorization(oauth, vault, fixture_lock)
        await authorization.remember(tokens)
        proxy = RemoteProxy(config, authorization)

        # Both sides use the official SDK: local MCP connection -> connector ->
        # real HTTP TCP -> application authorization middleware -> deployed tool.
        async with Client(proxy.server(), mode="legacy") as local_client:
            tools = await local_client.list_tools()
            assert "connection_status" in [tool.name for tool in tools.tools]
            result = await local_client.call_tool("connection_status", {})
            assert not result.is_error
            assert set(result.structured_content["capabilities"]) == {
                "mcp:read",
                "mcp:upload",
                "mcp:export",
            }
            assert result.structured_content["completeness"] == "complete"
            grant_id = uuid.UUID(result.structured_content["connection_id"])

            previous_refresh = vault.value.refresh_token
            authorization.tokens = None  # Simulate reconnect/expired in-memory access.
            refreshed = await local_client.call_tool("connection_status", {})
            assert not refreshed.is_error
            assert vault.value.refresh_token != previous_refresh

            async with sessions() as session:
                grant = await session.get(MCPGrantModel, grant_id)
                grant.revoked_at = datetime.now(UTC)
                grant.revocation_reason = "test_revocation"
                await session.commit()
            denied = await local_client.call_tool("connection_status", {})
            assert denied.is_error
            assert "connection_id" not in str(denied.structured_content)

        async with sessions() as session:
            hashes = list((await session.scalars(select(MCPTokenModel.token_hash))).all())
            assert tokens.access_token not in hashes and tokens.refresh_token not in hashes
            audits = list((await session.scalars(select(AuditLogModel))).all())
            assert any(
                row.action == "mcp.tool.connection_status" and row.result == "success"
                for row in audits
            )
            audit_text = json.dumps([row.metadata_json for row in audits], default=str)
            for secret in (
                code,
                tokens.access_token,
                tokens.refresh_token,
                vault.value.refresh_token,
            ):
                assert secret not in audit_text
        assert vault.value.authorized_until > time.time()


async def test_pdf_stage_workbook_download_and_verified_history_ack_over_tcp(backend_tcp, tmp_path):
    import hashlib
    import io

    from app.application.mcp.artifacts import MCPArtifactService
    from app.application.mcp.authorization import MCPAuthorizationService
    from app.infrastructure.database.models import (
        AgencyModel,
        ClientGroupModel,
        PassportExportHistoryModel,
        UntrustedUploadScanModel,
    )
    from app.infrastructure.repositories.passport_export_history_repository import (
        PassportExportHistoryRepository,
    )
    from app.infrastructure.security.upload_security import UploadSecurityService
    from app.infrastructure.storage.minio_repository import ObjectIntegrityMetadata
    from openpyxl import Workbook, load_workbook
    from pypdf import PdfWriter

    from gc_mcp_connector.artifacts import ArtifactClient

    config, sessions, attempt, code, user_id, app = backend_tcp

    class PrivateStorage:
        def __init__(self):
            self.data = {}

        async def put_transfer(self, source, *, key, size, sha256, media_type):
            assert key not in self.data
            source.seek(0)
            content = source.read()
            assert len(content) == size and hashlib.sha256(content).hexdigest() == sha256
            self.data[key] = (content, sha256, media_type)

        async def stat_file(self, key):
            content, checksum, media_type = self.data[key]
            return ObjectIntegrityMetadata(len(content), checksum, media_type)

        async def stream_file(self, key, *, start, expected_bytes, chunk_size):
            content = self.data[key][0]
            for offset in range(start, len(content), 113):
                yield content[offset : offset + 113]

    async def chunks(content):
        for offset in range(0, len(content), 113):
            yield content[offset : offset + 113]

    storage = PrivateStorage()
    app.state.mcp_artifact_storage = storage
    app.state.mcp_artifact_security = UploadSecurityService(
        settings=app.state.settings, session_factory=sessions, storage=storage
    )
    agency_id, group_id = uuid.uuid4(), uuid.uuid4()
    async with sessions() as session:
        session.add(
            AgencyModel(
                id=agency_id, name="Artifact TCP fixture", email="artifact-tcp@example.test"
            )
        )
        await session.flush()
        session.add(
            ClientGroupModel(
                id=group_id, agency_id=agency_id, name="Artifact TCP trip", token=uuid.uuid4().hex
            )
        )
        await session.commit()

    async with httpx2.AsyncClient(timeout=20, trust_env=False, follow_redirects=False) as http:
        oauth = OAuthClient(config, http)
        tokens = await oauth.exchange(code, attempt)
        authorization = Authorization(oauth, FixtureVault(), fixture_lock)
        await authorization.remember(tokens)
        transfers = ArtifactClient(config, authorization, http)
        document, buffer = PdfWriter(), io.BytesIO()
        document.add_blank_page(width=100, height=100)
        document.write(buffer)
        source = tmp_path / "explicit-provided.pdf"
        source.write_bytes(buffer.getvalue())
        staged = await transfers.upload_pdf(
            source,
            allowed_paths=frozenset({source.resolve()}),
            agency_id=agency_id,
            group_id=group_id,
        )
        assert staged.direction == "upload" and staged.byte_size == source.stat().st_size
        async with sessions() as session:
            scan = await session.scalar(select(UntrustedUploadScanModel))
            assert scan.scan_status == "clean" and scan.content_sha256 == staged.sha256
            principal = await MCPAuthorizationService(session, app.state.settings).verify_access(
                tokens.access_token
            )
            await session.commit()
            history = await PassportExportHistoryRepository(session).record(
                group_id=group_id,
                agency_id=agency_id,
                export_kind="passport_excel",
                export_mode="all",
                request_id=uuid.uuid4(),
                snapshot_submission_ids=[],
                exported_submission_ids=[],
                exported_people_snapshot=[],
                created_by_user_id=user_id,
                actor_email="connector-tcp@example.test",
            )
            book, output = Workbook(), io.BytesIO()
            book.active.append(["Passport Number", "Name"])
            book.save(output)
            artifact = await MCPArtifactService(
                session, app.state.settings, storage=storage
            ).prepare_export(
                principal,
                agency_id=agency_id,
                group_id=group_id,
                purpose="passport_excel",
                filename="fixture.xlsx",
                body=chunks(output.getvalue()),
                export_history_id=history.id,
            )
            history_id = history.id
            await session.commit()
        destination = tmp_path / "downloaded.xlsx"
        result = await transfers.download(artifact["artifact_id"], destination=destination)
        assert result.server_delivery_acknowledged
        assert result.file.sha256 == hashlib.sha256(output.getvalue()).hexdigest()
        assert list(load_workbook(destination).active.values) == [("Passport Number", "Name")]
        recovered = await transfers.acknowledge_existing(
            artifact["artifact_id"],
            path=destination,
            allowed_paths=frozenset({destination.resolve()}),
        )
        assert recovered.server_delivery_acknowledged
        async with sessions() as session:
            history = await session.get(PassportExportHistoryModel, history_id)
            assert history.status == "completed" and history.completed_at is not None
