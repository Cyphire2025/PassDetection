"""Synthetic SQLite All Data and export helper benchmark; never production capacity proof.

Creates a new output directory, seeds in batches of 500, then runs at most four
concurrent readers for 20-120 seconds. Uses the actual application read helper,
but does not include HTTP, authentication middleware, Redis or PostgreSQL.
"""

from __future__ import annotations

import argparse
import asyncio
import json
import os
import platform
import sqlite3
import sys
import time
import uuid
from datetime import UTC, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def identifier(kind: int, number: int = 0) -> uuid.UUID:
    return uuid.UUID(int=((0xF0000000 | kind) << 96) + number + 1)


async def measure(output: Path, records: int, seconds: int, concurrency: int) -> dict:
    # Isolate configuration before importing any application dependency. There
    # is no .env file in the newly created output directory, and no service call.
    os.chdir(output)
    os.environ.update(APP_ENV="development", APP_SECRET_KEY="synthetic-company-records-only",
                      POSTGRES_HOST="127.0.0.1", POSTGRES_PORT="1", POSTGRES_PASSWORD="synthetic-only",
                      S3_ACCESS_KEY_ID="synthetic-only", S3_SECRET_ACCESS_KEY="synthetic-only")
    sys.path.insert(0, str(ROOT))
    from cryptography.fernet import Fernet
    from sqlalchemy import event, insert, text, update
    from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
    from sqlalchemy.pool import AsyncAdaptedQueuePool

    from app.application.travellers.history_schedule import schedule_company_history_initialization
    from app.core.config.settings import Settings
    from app.core.config.traveller_reuse import TravellerReuseSettings
    from app.infrastructure.database.gc_mobile_models import ClientOrganizationModel
    from app.infrastructure.database.model_base import Base
    from app.infrastructure.database.models import (
        AgencyModel,
        ClientGroupModel,
        UserModel,
    )
    from app.infrastructure.database.traveller_reuse_models import (
        TravellerDocumentRevisionModel as Document,
    )
    from app.infrastructure.database.traveller_reuse_models import (
        TravellerParticipationModel as Participation,
    )
    from app.infrastructure.database.traveller_reuse_models import (
        TravellerPassportModel as Passport,
    )
    from app.infrastructure.database.traveller_reuse_models import TravellerProfileModel as Person
    from app.infrastructure.database.traveller_reuse_models import (
        TravellerReconciliationJobModel as HistoryJob,
    )
    from app.infrastructure.repositories.user_repository import UserRepository
    from app.infrastructure.security.traveller_reuse_crypto import TravellerReuseCrypto
    from app.presentation.api.v1.routes.company_people import people_page

    path = output / "synthetic-company-records.sqlite"
    engine = create_async_engine(f"sqlite+aiosqlite:///{path.as_posix()}", poolclass=AsyncAdaptedQueuePool,
                                 pool_size=concurrency, max_overflow=0)

    @event.listens_for(engine.sync_engine, "connect")
    def configure(connection, _record):
        cursor = connection.cursor()
        cursor.execute("PRAGMA foreign_keys=ON")
        cursor.execute("PRAGMA cache_size=-8192")
        cursor.close()

    agency, client, actor_id = identifier(1), identifier(2), identifier(8)
    crypto = TravellerReuseCrypto(TravellerReuseSettings(lookup_key="synthetic-all-data-only-" * 3, encryption_key=Fernet.generate_key().decode()))
    started = time.perf_counter()
    async with engine.begin() as connection:
        await connection.exec_driver_sql("PRAGMA journal_mode=WAL")
        await connection.run_sync(Base.metadata.create_all)
        await connection.execute(insert(AgencyModel), {"id": agency, "name": "Synthetic operator", "email": "operator@example.invalid"})
        await connection.execute(insert(ClientOrganizationModel), {"id": client, "agency_id": agency,
            "name": "Synthetic client", "normalized_name": "synthetic client"})
        await connection.execute(insert(ClientGroupModel), [{"id": identifier(3, index), "agency_id": agency,
            "company_id": client, "company_revision": 1, "name": f"Synthetic trip {index}", "token": f"synthetic-unusable-{index}",
            "status": ("active", "closed", "archived")[index % 3]} for index in range(100)])
        await connection.execute(insert(UserModel), {"id": actor_id, "agency_id": agency, "role": "agency_admin",
            "full_name": "Synthetic reviewer", "email": "reader@example.invalid", "hashed_password": "unused"})
        for start in range(0, records, 500):
            if start % 25000 == 0:
                print(json.dumps({"phase": "seeding", "completed": start, "target": records}), flush=True)
            numbers = range(start, min(start + 500, records))
            common = {"agency_id": agency, "company_id": client}
            await connection.execute(insert(Person), [{**common, "id": identifier(4, index), "status": "active",
                "display_name_ciphertext": crypto.encrypt(f"Synthetic Traveller {index:06d}"),
                "name_fingerprint": crypto.scoped_digest("name", agency, client, f"synthetic traveller {index:06d}"),
                "fingerprint_key_version": crypto.key_version,
                "created_at": datetime(2026, 1, 1, tzinfo=UTC) + timedelta(seconds=index)} for index in numbers])
            await connection.execute(insert(Passport), [{**common, "id": identifier(5, index), "traveller_id": identifier(4, index),
                "issuing_country": "IND", "document_type": "PASSPORT", "number_ciphertext": crypto.encrypt(f"T{index:08d}"),
                "number_fingerprint": crypto.passport_fingerprint(agency, client, f"T{index:08d}", "IND", "PASSPORT"),
                "fingerprint_key_version": crypto.key_version, "status": "accepted", "is_default": True} for index in numbers])
            await connection.execute(insert(Document), [{**common, "id": identifier(6, index), "traveller_id": identifier(4, index),
                "passport_id": identifier(5, index), "revision_number": 1, "status": "accepted", "accepted_at": datetime.now(UTC),
                "fields_ciphertext": crypto.encrypt(json.dumps({"passport_number": f"T{index:08d}", "given_names": "Synthetic",
                    "surname": f"Traveller {index:06d}", "date_of_birth": "1990-01-01", "date_of_expiry": "2035-01-01"})),
                "idempotency_key": f"synthetic:{index}"} for index in numbers])
            await connection.execute(insert(Participation), [{**common, "id": identifier(7, index), "traveller_id": identifier(4, index),
                "group_id": identifier(3, index % 100), "passport_id": identifier(5, index), "document_revision_id": identifier(6, index),
                "source_kind": "synthetic", "source_row_key": str(index), "profile_revision": 1, "contact_revision": 1,
                "status": "accepted", "question_version": "synthetic", "idempotency_key": f"synthetic:{index}"} for index in numbers])
        await connection.execute(text("ANALYZE"))
    seed_seconds = time.perf_counter() - started
    settings = Settings(_env_file=None, app_env="development")
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    async with sessions() as session, session.begin():
        actor = await UserRepository(session).get_by_id(actor_id)
        company = await session.get(ClientOrganizationModel, client)
        await schedule_company_history_initialization(session, company=company)
        await session.execute(update(HistoryJob).values(status="succeeded", last_error_code=None))
    statements: dict[str, tuple] = {}

    @event.listens_for(engine.sync_engine, "before_cursor_execute")
    def capture(_connection, _cursor, statement, parameters, _context, _many):
        nonlocal query_total
        if statement.startswith("SELECT"):
            query_total += 1
        if statement.startswith("SELECT") and len(statements) < 16:
            statements.setdefault(statement, parameters)

    samples: dict[str, list[float]] = {"first_page": [], "cursor_page": [], "name_search": []}
    failures = 0
    query_total = 0
    maximum_returned = 0
    deadline = time.perf_counter() + seconds

    async def reader(worker: int):
        nonlocal failures, maximum_returned
        cursor = None
        iteration = worker
        while time.perf_counter() < deadline:
            kind = "name_search" if iteration % 5 == 0 else "cursor_page" if cursor else "first_page"
            started = time.perf_counter()
            try:
                async with sessions() as session:
                    result = await people_page(company_id=client, agency_id=agency,
                        search="Synthetic Traveller 000001" if kind == "name_search" else "",
                        cursor=cursor if kind == "cursor_page" else None, page_size=25,
                        user=actor, session=session, settings=settings, crypto=crypto)
                maximum_returned = max(maximum_returned, len(result.items))
                assert len(result.items) <= 25
                if kind == "name_search":
                    assert len(result.items) == 1 and result.items[0].id == identifier(4, 1)
                else:
                    assert len(result.items) == 25
                    cursor = result.next_cursor if kind == "first_page" else None
                # Bound retained timings independently from throughput.
                if len(samples[kind]) < 10000:
                    samples[kind].append((time.perf_counter() - started) * 1000)
            except Exception:
                failures += 1
                raise
            iteration += 1

    load_started = time.perf_counter()
    await asyncio.gather(*(reader(index) for index in range(concurrency)))
    elapsed = time.perf_counter() - load_started
    print(json.dumps({"phase": "reads_finished", "errors": failures, "elapsed_seconds": round(elapsed, 3)}), flush=True)
    (output / "read_samples.json").write_text(json.dumps({"samples_ms": samples, "queries": query_total, "errors": failures}), encoding="utf-8")
    read_query_total = query_total
    before_query_count = query_total
    async with sessions() as session:
        await people_page(company_id=client, agency_id=agency, search="", cursor=None, page_size=25,
            user=actor, session=session, settings=settings, crypto=crypto)
    first_page_query_count = query_total - before_query_count
    export = await measure_export(output, sessions, actor, client, agency, crypto, records)
    plans = []
    async with engine.connect() as connection:
        for statement, parameters in statements.items():
            rows = (await connection.exec_driver_sql("EXPLAIN QUERY PLAN " + statement, parameters)).all()
            plans.append({"sql": statement, "plan": [row[-1] for row in rows]})
    await engine.dispose()

    def stats(values: list[float]) -> dict:
        ordered = sorted(values)
        return {"sampled_operations": len(values), **{name: round(ordered[min(len(ordered) - 1, int((len(ordered) - 1) * quantile))], 3)
            for name, quantile in (("p50_ms", .5), ("p95_ms", .95), ("p99_ms", .99))}}

    peak = None
    if os.name == "nt":
        import subprocess
        peak = int(subprocess.check_output(["powershell", "-NoProfile", "-Command",
            f"(Get-Process -Id {os.getpid()}).PeakWorkingSet64"], text=True, timeout=10).strip())
    return {"dataset": {"canonical_people": records, "accepted_document_versions": records, "client_companies": 1, "assigned_groups": 100},
        "runtime": {"python": platform.python_version(), "sqlite": sqlite3.sqlite_version, "platform": platform.system()},
        "seed_batch_size": 500, "seed_seconds": round(seed_seconds, 3), "concurrent_readers": concurrency,
        "elapsed_seconds": round(elapsed, 3), "errors": failures, "max_records_returned": maximum_returned,
        "database_bytes": path.stat().st_size, "peak_working_set_bytes": peak,
        "first_page_query_count": first_page_query_count, "read_query_total": read_query_total, "export": export,
        "budgets": {"read_p95_ms": 250, "read_p99_ms": 500, "page_size": 25, "export_page_people": 250, "export_volume_bytes": 64 * 1024 * 1024, "export_temp_xml_bytes": 256 * 1024 * 1024},
        "read_budgets_met": all(stats(values)["p95_ms"] <= 250 and stats(values)["p99_ms"] <= 500 for values in samples.values()),
        "latencies": {name: stats(values) for name, values in samples.items()}, "query_plans": plans,
        "limitations": ["Synthetic SQLite application-helper benchmark, not PostgreSQL or HTTP capacity",
            "No OTP, document bytes, authentication middleware or network", "No million-user or sustained-production claim",
            "Exact name search uses tenant HMAC index; no fuzzy or substring search",
            "One client with 100 groups and one version per person; sparse-client and long per-person histories not load-tested",
            "Excel tested at full dataset size; ZIP byte boundaries and provider failures covered by focused tests only"]}


async def measure_export(output, sessions, actor, company, agency, crypto, records):
    from openpyxl import load_workbook

    from app.infrastructure.travellers import company_export_worker as worker
    from app.infrastructure.travellers.company_export_render import WorkbookAssembly
    peak_xml = 0
    class MeasuredWorkbook(WorkbookAssembly):
        def append(self, rows):
            nonlocal peak_xml
            super().append(rows)
            peak_xml = max(peak_xml, Path(self.sheet._writer.out).stat().st_size)
    worker.WorkbookAssembly = MeasuredWorkbook
    from sqlalchemy import select

    from app.application.companies.data_exports import ExportRequest, request_export
    from app.infrastructure.database.company_data_export_models import (
        CompanyDataExportJobModel as Job,
    )
    from app.infrastructure.database.company_data_export_models import (
        CompanyDataExportPartModel as Part,
    )
    from app.infrastructure.travellers.company_export_worker import run_company_export_tick

    class PrivateDiskStorage:
        def __init__(self):
            self.root = output / "private-export-parts"
            self.root.mkdir(mode=0o700)
            self.fail_once = True
            self.max_upload = 0
        def path(self, key):
            import hashlib
            return self.root / hashlib.sha256(key.encode()).hexdigest()
        async def upload_file(self, content, key, content_type):
            self.max_upload = max(self.max_upload, len(content))
            self.path(key).write_bytes(content)
            if self.fail_once:
                self.fail_once = False
                raise TimeoutError("Synthetic acknowledgement loss")
            return key
        async def stream_file(self, key, *, start, expected_bytes, chunk_size=65536):
            path = self.path(key)
            assert path.stat().st_size == expected_bytes
            with path.open("rb") as source:
                source.seek(start)
                while chunk := source.read(chunk_size):
                    yield chunk
    storage = PrivateDiskStorage()
    async with sessions() as session, session.begin():
        job = await request_export(session, user=actor, agency_id=agency, company_id=company,
            body=ExportRequest(kind="excel", idempotency_key="synthetic-full-capacity"))
        job_id = job.id
    started = time.perf_counter()
    pages = 0
    assert await run_company_export_tick(session_factory=sessions, storage=storage, crypto=crypto) == 0
    async with sessions() as session, session.begin():
        job = await session.get(Job, job_id)
        assert job.processed_count == 0 and job.part_count == 0 and job.cursor is None
        assert await session.scalar(select(Part.id).where(Part.job_id == job_id)) is not None
        job.available_at = job.lease_expires_at = datetime.now(UTC) - timedelta(seconds=1)
    while pages < records // 250 + 10:
        assert await run_company_export_tick(session_factory=sessions, storage=storage, crypto=crypto) == 1
        pages += 1
        async with sessions() as session:
            job = await session.get(Job, job_id)
            if job.status == "succeeded":
                part = await session.scalar(select(Part).where(Part.job_id == job_id, Part.part_kind == "workbook"))
                break
    assert job.status == "succeeded" and job.processed_count == records
    with storage.path(part.object_key).open("rb") as stream:
        workbook = load_workbook(stream, read_only=True)
        count = sum(1 for _ in workbook.active.values) - 1
        workbook.close()
    assert count == records
    return {"elapsed_seconds": round(time.perf_counter() - started, 3), "people": count,
        "worker_ticks": pages, "resumed_after_acknowledgement_loss": True,
        "workbook_bytes": part.byte_size, "maximum_stored_part_bytes": storage.max_upload,
        "peak_temporary_xml_bytes": peak_xml,
        "private_part_disk_bytes": sum(path.stat().st_size for path in storage.root.iterdir())}


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--output", type=Path, required=True)
    parser.add_argument("--records", type=int, default=100000)
    parser.add_argument("--seconds", type=int, default=60)
    parser.add_argument("--readers", type=int, default=4, choices=range(1, 5))
    args = parser.parse_args()
    if not 50000 <= args.records <= 100000 or not 20 <= args.seconds <= 120:
        parser.error("Use 50000-100000 records and 20-120 seconds.")
    output = args.output.resolve()
    output.mkdir(parents=True, exist_ok=False)
    result = asyncio.run(measure(output, args.records, args.seconds, args.readers))
    (output / "result.json").write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps({key: value for key, value in result.items() if key != "query_plans"}, indent=2))


if __name__ == "__main__":
    main()
