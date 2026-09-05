"""Synthetic local comparison of passport export and page allocation costs.

Run from backend with its Python environment. Uses only an in-memory SQLite
database and invented data; it never connects to application infrastructure.
Baseline code is read from the selected local Git revision into temporary files.
This is a diagnostic benchmark, not a production throughput claim.
"""

from __future__ import annotations

import argparse
import asyncio
import gc
import importlib.util
import io
import json
import platform
import statistics
import subprocess
import sys
import tempfile
import time
import tracemalloc
import uuid
from pathlib import Path

from openpyxl import load_workbook
from sqlalchemy import select
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

from app.application.dtos.passport_dtos import passport_submission_output_from_entity
from app.application.use_cases.passports.submission_view import build_submission_view
from app.domain.entities.entities import PassportSubmission, User, UserRole
from app.infrastructure.database.models import ClientGroupModel, PassportSubmissionModel
from app.infrastructure.export.passport_excel_exporter import PassportExcelExporter
from app.infrastructure.repositories.passport_submission_repository import (
    PassportSubmissionRepository,
)
from app.infrastructure.repositories.passport_submission_view_repository import (
    PassportSubmissionViewRepository,
)


def baseline_module(root: Path, revision: str, relative: str, name: str):
    source = subprocess.check_output(["git", "show", f"{revision}:{relative}"], cwd=root)
    path = Path(tempfile.gettempdir()) / f"{name}_{uuid.uuid4().hex}.py"
    try:
        path.write_bytes(source)
        spec = importlib.util.spec_from_file_location(name, path)
        module = importlib.util.module_from_spec(spec)
        sys.modules[name] = module
        spec.loader.exec_module(module)
        return module
    finally:
        path.unlink(missing_ok=True)


def workbook_contract(content: bytes):
    sheet = load_workbook(io.BytesIO(content)).active
    # Generated-at cell intentionally varies; values/styles/table extents must not.
    return (
        [
            (cell.coordinate, cell.value, cell.style_id, cell.number_format)
            for row in sheet
            for cell in row
            if cell.coordinate != "A2"
        ],
        [table.ref for table in sheet.tables.values()],
    )


def export_benchmark(size, baseline, repeats):
    group_id, agency_id = uuid.uuid4(), uuid.uuid4()
    rows = []
    for index in range(size):
        submission = PassportSubmission.create(
            group_id=group_id,
            agency_id=agency_id,
            client_name=f"Synthetic {index:05}",
            client_email=None,
            image_s3_key="synthetic/front.jpg",
        )
        submission.submit_client_review(
            {
                "given_names": f"Synthetic {index:05}",
                "surname": "Traveller",
                "passport_number": f"T{index:07}",
                "date_of_birth": "1990-01-01",
                "nationality": "IND",
            },
            client_email="synthetic@example.test",
            client_phone=f"+9199{index:08}",
            departure_city="Delhi",
            nearest_domestic_airport="Delhi",
        )
        rows.append(submission)
    kwargs = dict(
        group_name="Synthetic benchmark",
        group_by_field="zone_name",
        additional_fields=[{"key": "zone_name", "label": "Zone"}],
        zone_names={row.id: str(index % 3) for index, row in enumerate(rows)},
        pending_rows=[{"GIVEN NAME": f"Pending {i}", "Zone": str(i % 3)} for i in range(10)],
    )
    timings, contracts = {}, []
    for name, exporter in [
        ("baseline", baseline.PassportExcelExporter()),
        ("current", PassportExcelExporter()),
    ]:
        samples = []
        for _ in range(repeats):
            gc.collect()
            started = time.perf_counter()
            content = exporter.export_group(rows, **kwargs)
            samples.append(time.perf_counter() - started)
        timings[f"{name}_seconds_median"] = round(statistics.median(samples), 4)
        contracts.append(workbook_contract(content))
    return {
        "submissions": size,
        "pending_rows": 10,
        **timings,
        "workbook_contract_equal": contracts[0] == contracts[1],
    }


async def view_benchmark(size, baseline):
    engine = create_async_engine("sqlite+aiosqlite:///:memory:")
    factory = async_sessionmaker(engine, expire_on_commit=False)
    user = User(
        id=uuid.uuid4(),
        agency_id=uuid.uuid4(),
        full_name="Synthetic admin",
        email="synthetic@example.test",
        hashed_password="unused",
        role=UserRole.AGENCY_ADMIN,
    )
    group_id = uuid.uuid4()
    try:
        async with engine.begin() as connection:
            await connection.run_sync(ClientGroupModel.__table__.create)
            await connection.run_sync(PassportSubmissionModel.__table__.create)
        async with factory() as session:
            session.add(
                ClientGroupModel(
                    id=group_id, agency_id=user.agency_id, name="Synthetic", token=uuid.uuid4().hex
                )
            )
            session.add_all(
                [
                    PassportSubmissionModel(
                        id=uuid.uuid4(),
                        agency_id=user.agency_id,
                        group_id=group_id,
                        client_name=f"Synthetic {index:05}",
                        image_s3_key="synthetic/front.jpg",
                        status="submitted",
                        extracted_fields={
                            "given_names": f"Synthetic {index:05}",
                            "surname": "Traveller",
                            "passport_number": f"T{index:07}",
                        },
                        mrz_raw="S" * 4096,
                        staff_metadata={"synthetic_trace": "D" * 4096},
                    )
                    for index in range(size)
                ]
            )
            await session.commit()
        stats, ids = {}, []
        for mode in ("baseline", "current"):
            gc.collect()
            tracemalloc.start()
            started = time.perf_counter()
            async with factory() as session:
                if mode == "baseline":
                    result = await session.execute(
                        select(PassportSubmissionModel).where(
                            PassportSubmissionModel.group_id == group_id
                        )
                    )
                    full_rows = [
                        passport_submission_output_from_entity(
                            PassportSubmissionRepository._to_entity(model)
                        )
                        for model in result.scalars()
                    ]
                    view = baseline.build_submission_view(
                        full_rows,
                        submission_filter="all",
                        sort_by="name",
                        sort_order="asc",
                        search=None,
                        page=1,
                        page_size=50,
                    )
                    hydrated = len(full_rows)
                else:
                    repository = PassportSubmissionViewRepository(session)
                    projection = await repository.projection(
                        group_id=group_id, user=user, include_deleted=False
                    )
                    view = build_submission_view(
                        projection,
                        submission_filter="all",
                        sort_by="name",
                        sort_order="asc",
                        search=None,
                        page=1,
                        page_size=50,
                    )
                    details = await repository.page_details(
                        submission_ids=[item.submission.id for item in view.items],
                        group_id=group_id,
                        user=user,
                    )
                    hydrated = len(details)
                ids.append(view.ordered_submission_ids)
                peak = tracemalloc.get_traced_memory()[1]
            stats[mode] = {
                "elapsed_with_tracemalloc_seconds": round(time.perf_counter() - started, 4),
                "peak_python_mib": round(peak / 1024 / 1024, 3),
                "full_dto_rows": hydrated,
            }
            tracemalloc.stop()
        return {
            "rows": size,
            "page_size": 50,
            "synthetic_non_projection_bytes_per_row": 8192,
            **stats,
            "selection_order_equal": ids[0] == ids[1],
        }
    finally:
        await engine.dispose()


async def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--baseline", default="HEAD")
    parser.add_argument("--repeats", type=int, default=3)
    parser.add_argument("--output", type=Path, required=True)
    args = parser.parse_args()
    root = Path(__file__).resolve().parents[2]
    original_export = baseline_module(
        root,
        args.baseline,
        "backend/app/infrastructure/export/passport_excel_exporter.py",
        "baseline_passport_export",
    )
    original_view = baseline_module(
        root,
        args.baseline,
        "backend/app/application/use_cases/passports/submission_view.py",
        "baseline_passport_view",
    )
    results = {
        "python": platform.python_version(),
        "platform": platform.platform(),
        "baseline_revision": subprocess.check_output(
            ["git", "rev-parse", args.baseline], cwd=root, text=True
        ).strip(),
        "repeats_export": args.repeats,
        "limits": "Synthetic local CPU/Python allocation measurements; SQLite query costs do not predict production PostgreSQL/network load. Current view still scans all authorized identity projections. 3000-row export is direct exporter stress; route limit is 1500 including pending rows.",
        "export": [],
        "view": [],
    }
    for size in (800, 1500, 3000):
        results["export"].append(export_benchmark(size, original_export, args.repeats))
        results["view"].append(await view_benchmark(size, original_view))
        args.output.write_text(json.dumps(results, indent=2) + "\n", encoding="utf-8")
        print(
            json.dumps({"export": results["export"][-1], "view": results["view"][-1]}), flush=True
        )


if __name__ == "__main__":
    asyncio.run(main())
