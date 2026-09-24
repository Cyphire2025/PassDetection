from __future__ import annotations

import io
import uuid
from unittest.mock import AsyncMock

import pytest
from openpyxl import load_workbook

from app.domain.entities.entities import PassportSubmission
from app.infrastructure.export.passport_excel_exporter import PassportExcelExporter
from app.presentation.api.v1.routes.passport_routes.ecr_export_support import (
    export_passport_ecr_results,
)


def submission(group_id, name="Sample traveller"):
    item = PassportSubmission.create(group_id, uuid.uuid4(), name, None, "front.jpg")
    item.confirmed_fields = {"given_names": name, "ECR": "ECR", "ecr": "ECR"}
    return item


def workbook(items, details=None, results=None, **options):
    content = PassportExcelExporter().export_group(
        items, group_name="Sample group", group_details=details, ecr_results=results, **options,
    )
    return load_workbook(io.BytesIO(content)).active


def result_cells(sheet):
    headers = [cell.value for cell in sheet[4]]
    result_column = headers.index("ECR") + 1
    name_column = headers.index("GIVEN NAME") + 1
    return {
        sheet.cell(row, name_column).value: sheet.cell(row, result_column)
        for row in range(5, sheet.max_row + 1)
    }


@pytest.mark.parametrize("details", [None, {}, {"passport_ecr_enabled": False}])
def test_disabled_or_legacy_group_exports_have_no_ecr_column(details):
    group_id = uuid.uuid4()
    item = submission(group_id)
    sheet = workbook([item], {group_id: details} if details is not None else None, {item.id: "ECR"})
    assert "ECR" not in [cell.value for cell in sheet[4]]


@pytest.mark.parametrize(("verdict", "expected", "color"), [
    ("ECR", "ECR", "00FF0000"), ("NA", "NA", "00000000"),
    ("REVIEW", "REVIEW", "00B45309"), ("ERROR", "ERROR", "00B45309"),
    ("PENDING", "PENDING", "00B45309"), ("NO_BACK", "REVIEW", "00B45309"),
])
def test_enabled_group_exports_trusted_verdict_with_standalone_ecr_colors(verdict, expected, color):
    group_id = uuid.uuid4()
    item = submission(group_id)
    sheet = workbook([item], {group_id: {"passport_ecr_enabled": True}}, {item.id: verdict})
    cell = result_cells(sheet)["SAMPLE TRAVELLER"]
    assert cell.value == expected
    assert cell.font.color.rgb == color
    assert cell.font.bold == (verdict == "ECR")


def test_missing_check_is_pending_and_does_not_use_untrusted_passport_fields():
    group_id = uuid.uuid4()
    sheet = workbook([submission(group_id)], {group_id: {"passport_ecr_enabled": True}})
    assert result_cells(sheet)["SAMPLE TRAVELLER"].value == "PENDING"


def test_combined_export_keeps_disabled_and_not_submitted_rows_blank():
    enabled, disabled = uuid.uuid4(), uuid.uuid4()
    checked, unchecked = submission(enabled, "Checked"), submission(disabled, "Unchecked")
    sheet = workbook(
        [checked, unchecked],
        {enabled: {"passport_ecr_enabled": True}, disabled: {"passport_ecr_enabled": False}},
        {checked.id: "ECR", unchecked.id: "ECR"},
        pending_rows=[{"GIVEN NAME": "Not submitted"}],
    )
    cells = result_cells(sheet)
    assert cells["CHECKED"].value == "ECR"
    assert cells["UNCHECKED"].value is None
    assert cells["Not submitted"].value is None


async def test_disabled_export_does_not_query_check_results():
    group_id, agency_id = uuid.uuid4(), uuid.uuid4()
    session = AsyncMock()
    assert await export_passport_ecr_results(
        session, [submission(group_id)], agency_id=agency_id,
        group_details={group_id: {"passport_ecr_enabled": False}},
    ) == {}
    session.execute.assert_not_awaited()


async def test_export_looks_up_only_supplied_opted_in_submissions_with_agency_scope(monkeypatch):
    from app.infrastructure.ecr import passport_runtime

    enabled, disabled, unknown = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    items = [submission(enabled), submission(disabled), submission(unknown)]
    agency_id = uuid.uuid4()
    session = AsyncMock()
    lookup = AsyncMock(return_value={items[0].id: "ECR"})
    monkeypatch.setattr(passport_runtime, "passport_ecr_results", lookup)
    results = await export_passport_ecr_results(
        session, items, agency_id=agency_id,
        group_details={enabled: {"passport_ecr_enabled": True}, disabled: {"passport_ecr_enabled": False}},
    )
    assert results == {items[0].id: "ECR"}
    lookup.assert_awaited_once_with(
        session, [items[0].id], agency_id=agency_id,
        expected_source_keys={items[0].id: items[0].passport_back_s3_key},
    )
