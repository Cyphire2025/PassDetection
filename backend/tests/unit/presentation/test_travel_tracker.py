"""Canonical roster, permissions, atomic marking, and spreadsheet round trips."""

from __future__ import annotations

import io
import uuid
from dataclasses import replace
from types import SimpleNamespace

import pytest
from openpyxl import load_workbook
from pydantic import ValidationError
from sqlalchemy import select
from sqlalchemy.exc import IntegrityError

from app.domain.entities.entities import UserRole
from app.domain.exceptions.travel_tracker import TravelTrackerError
from app.infrastructure.database.models import (
    AgencyModel,
    AuditLogModel,
    ClientGroupModel,
    PassportRosterResolutionModel,
    PassportSubmissionModel,
    UserModel,
)
from app.infrastructure.database.travel_tracker_model import TravelTrackerModel
from app.infrastructure.repositories.user_repository import UserRepository
from app.infrastructure.travel_tracker.service import TravelTrackerService
from app.presentation.api.v1.schemas.travel_tracker_schemas import TrackerMarkRequest


@pytest.fixture
async def tracker(db_session):
    agency, foreign_agency, actor_id = uuid.uuid4(), uuid.uuid4(), uuid.uuid4()
    db_session.add_all(
        [
            AgencyModel(id=agency, name="Synthetic", email=f"{agency}@example.test"),
            AgencyModel(id=foreign_agency, name="Foreign", email=f"{foreign_agency}@example.test"),
        ]
    )
    await db_session.flush()
    actor = UserModel(
        id=actor_id,
        agency_id=agency,
        full_name="Office user",
        role="agency_admin",
        email=f"{actor_id}@example.test",
        hashed_password="unused",
    )
    db_session.add(actor)
    await db_session.flush()
    group = ClientGroupModel(
        id=uuid.uuid4(),
        agency_id=agency,
        created_by_user_id=actor_id,
        name="Canonical group",
        token=uuid.uuid4().hex,
        custom_questions=[
            {
                "id": "meal",
                "label": "Meal",
                "enabled": True,
                "required": False,
                "options": ["Veg", "Non veg"],
            }
        ],
        custom_details=[{"id": "code", "label": "Employee ID", "enabled": True, "required": False}],
    )
    imported = ClientGroupModel(
        id=uuid.uuid4(),
        agency_id=agency,
        name="Imported group",
        token=uuid.uuid4().hex,
        status="closed",
        import_only=True,
    )
    foreign = ClientGroupModel(
        id=uuid.uuid4(), agency_id=foreign_agency, name="Foreign group", token=uuid.uuid4().hex
    )
    db_session.add_all([group, imported, foreign])
    await db_session.flush()
    passengers = [
        PassportSubmissionModel(
            id=uuid.uuid4(),
            group_id=group.id,
            agency_id=agency,
            client_name=name,
            status=state,
            image_s3_key="excel-imports/synthetic" if index == 1 else "synthetic/front.jpg",
            confirmed_fields={
                "given_names": name,
                "passport_number": f"P{index}",
                "staff_code": "E42",
            },
            extracted_fields={"nationality": "IND", "date_of_birth": "1990-01-02"},
            staff_metadata={"zone_name": "West", "designation": "Engineer"},
            custom_answers=[{"question_id": "meal", "label": "Meal", "value": "Veg"}],
            custom_detail_answers=[{"detail_id": "code", "label": "Employee ID", "value": "EMP42"}],
        )
        for index, (name, state) in enumerate(
            [
                ("Asha Rao", "ai_approved"),
                ("Imported pending", "pending_upload"),
                ("Needs Review", "needs_review"),
                ("Failed extraction", "failed"),
                ("Same Name", "submitted"),
                ("Same Name", "processing"),
            ]
        )
    ]
    extra = PassportSubmissionModel(
        id=uuid.uuid4(),
        group_id=imported.id,
        agency_id=agency,
        client_name="Imported only",
        status="needs_review",
        image_s3_key="excel-imports/synthetic",
    )
    foreign_passenger = PassportSubmissionModel(
        id=uuid.uuid4(),
        group_id=foreign.id,
        agency_id=foreign_agency,
        client_name="Foreign only",
        status="submitted",
        image_s3_key="synthetic",
    )
    db_session.add_all([*passengers, extra, foreign_passenger])
    await db_session.commit()
    user = UserRepository._to_entity(actor)
    return SimpleNamespace(
        session=db_session,
        user=user,
        actor=actor,
        group=group,
        passengers=passengers,
        imported=imported,
        foreign=foreign,
        foreign_passenger=foreign_passenger,
        service=TravelTrackerService(db_session, user),
    )


async def test_complete_canonical_roster_includes_all_states_and_import_groups(tracker):
    groups = await tracker.service.list_groups()
    assert {group.id for group in groups.groups} == {tracker.group.id, tracker.imported.id}
    assert next(group for group in groups.groups if group.id == tracker.group.id).total == 6
    page = await tracker.service.workspace(tracker.group.id, page_size=2)
    assert page.counts.total == page.total == 6 and len(page.passengers) == 2
    assert page.counts.marked == 0 and page.counts.pending == 6
    whole = await tracker.service.workspace(tracker.group.id)
    assert {row.id for row in whole.passengers} == {row.id for row in tracker.passengers}
    assert {row.submission_status for row in whole.passengers} >= {
        "pending_upload",
        "failed",
        "needs_review",
    }


async def test_explicit_roster_removal_is_excluded_from_reads_and_marks(tracker):
    removed = tracker.passengers[0]
    tracker.session.add(
        PassportRosterResolutionModel(
            id=uuid.uuid4(),
            agency_id=tracker.group.agency_id,
            client_group_id=tracker.group.id,
            submission_id=removed.id,
            resolution_type="rejected",
            status="active",
        )
    )
    await tracker.session.commit()
    workspace = await tracker.service.workspace(tracker.group.id)
    assert workspace.total == workspace.group.total == 5
    assert removed.id not in {row.id for row in workspace.passengers}
    with pytest.raises(TravelTrackerError) as error:
        await tracker.service.mark(
            tracker.group.id,
            TrackerMarkRequest(track="visa", marked=True, passenger_ids=[removed.id]),
        )
    assert error.value.status_code == 404


async def test_visa_flight_independent_idempotent_attributed_and_audited(tracker):
    ids = [row.id for row in tracker.passengers[:3]]
    before = [(row.status, row.confirmed_fields.copy()) for row in tracker.passengers]
    result = await tracker.service.mark(
        tracker.group.id,
        TrackerMarkRequest(track="visa", marked=True, passenger_ids=[*ids, ids[0]]),
    )
    assert result.updated_count == 3 and result.unchanged_count == 0 and result.counts.marked == 3
    assert (
        await tracker.service.mark(
            tracker.group.id, TrackerMarkRequest(track="visa", marked=True, passenger_ids=ids)
        )
    ).updated_count == 0
    await tracker.service.mark(
        tracker.group.id, TrackerMarkRequest(track="flight", marked=True, passenger_ids=[ids[0]])
    )
    await tracker.service.mark(
        tracker.group.id, TrackerMarkRequest(track="visa", marked=False, passenger_ids=[ids[0]])
    )
    marks = {
        row.passenger_id: row
        for row in (await tracker.session.scalars(select(TravelTrackerModel))).all()
    }
    assert marks[ids[0]].flight_booked and not marks[ids[0]].visa_applied
    assert marks[ids[1]].visa_applied and not marks[ids[1]].flight_booked
    assert marks[ids[0]].visa_updated_by == marks[ids[0]].flight_updated_by == tracker.user.id
    assert marks[ids[0]].visa_updated_at and marks[ids[0]].flight_updated_at
    logs = list(
        (
            await tracker.session.scalars(
                select(AuditLogModel).where(AuditLogModel.action == "travel_tracker.marks_changed")
            )
        ).all()
    )
    assert len(logs) == 3 and logs[0].user_id == tracker.user.id
    for passenger, original in zip(tracker.passengers, before, strict=True):
        await tracker.session.refresh(passenger)
        assert (passenger.status, passenger.confirmed_fields) == original


@pytest.mark.parametrize("invalid", ["foreign", "missing"])
async def test_bulk_invalid_ids_are_atomic(tracker, invalid):
    invalid_id = tracker.foreign_passenger.id if invalid == "foreign" else uuid.uuid4()
    with pytest.raises(TravelTrackerError) as error:
        await tracker.service.mark(
            tracker.group.id,
            TrackerMarkRequest(
                track="visa", marked=True, passenger_ids=[tracker.passengers[0].id, invalid_id]
            ),
        )
    assert error.value.status_code == 404
    await tracker.session.rollback()
    assert not (await tracker.session.scalars(select(TravelTrackerModel))).all()
    assert not (await tracker.session.scalars(select(AuditLogModel))).all()


async def test_filtered_selection_uses_same_search_and_rejects_stale_count(tracker):
    filtered = await tracker.service.workspace(
        tracker.group.id, status="pending", search="P2", page_size=1
    )
    assert filtered.total == 1 and filtered.counts.total == 6
    with pytest.raises(TravelTrackerError) as error:
        await tracker.service.mark(
            tracker.group.id,
            TrackerMarkRequest(
                track="visa",
                marked=True,
                selection={"status": "pending", "search": "P2"},
                expected_count=2,
            ),
        )
    assert error.value.status_code == 409
    assert not (await tracker.session.scalars(select(TravelTrackerModel))).all()
    result = await tracker.service.mark(
        tracker.group.id,
        TrackerMarkRequest(
            track="visa",
            marked=True,
            selection={"status": "pending", "search": "P2"},
            expected_count=1,
        ),
    )
    assert result.passenger_ids == [tracker.passengers[2].id]
    marked = await tracker.service.workspace(tracker.group.id, status="marked")
    assert marked.total == 1 and marked.counts.pending == 5
    pending = await tracker.service.workspace(tracker.group.id, status="pending", search="Same")
    assert pending.total == 2 and pending.counts.marked == 1


@pytest.mark.parametrize(
    "restriction",
    [
        "other_agency",
        "staff_unassigned",
        "coordinator",
        "client_manager",
        "disabled",
        "role_changed",
    ],
)
async def test_write_authorization_revalidated(tracker, restriction):
    user = tracker.user
    if restriction == "other_agency":
        tracker.actor.agency_id = tracker.foreign.agency_id
        user = UserRepository._to_entity(tracker.actor)
    elif restriction == "staff_unassigned":
        tracker.actor.role, tracker.group.created_by_user_id = "agency_staff", None
        user = UserRepository._to_entity(tracker.actor)
    elif restriction in {"coordinator", "client_manager"}:
        user = replace(
            user,
            role=UserRole.AGENCY_COORDINATOR
            if restriction == "coordinator"
            else UserRole.CLIENT_MANAGER,
        )
    elif restriction == "disabled":
        tracker.actor.is_active = False
    else:
        tracker.actor.role = "agency_coordinator"
    await tracker.session.commit()
    with pytest.raises(TravelTrackerError) as error:
        await TravelTrackerService(tracker.session, user).mark(
            tracker.group.id,
            TrackerMarkRequest(track="visa", marked=True, passenger_ids=[tracker.passengers[0].id]),
        )
    assert error.value.status_code in {403, 404}
    assert not (await tracker.session.scalars(select(TravelTrackerModel))).all()


async def test_credential_allowlist_scopes_pagination_and_objects(tracker):
    service = TravelTrackerService(
        tracker.session, tracker.user, allowed_group_ids=[tracker.imported.id]
    )
    groups = await service.list_groups(page_size=1)
    assert groups.total == 1 and groups.groups[0].id == tracker.imported.id
    with pytest.raises(TravelTrackerError) as error:
        await service.workspace(tracker.group.id)
    assert error.value.status_code == 404
    assert (
        await TravelTrackerService(
            tracker.session, tracker.user, allowed_group_ids=[]
        ).list_groups()
    ).total == 0


@pytest.mark.parametrize("lifecycle", ["archived", "deleted"])
async def test_removed_groups_not_mutable_even_global_superadmin(tracker, lifecycle):
    tracker.group.status = lifecycle
    tracker.actor.role, tracker.actor.agency_id = "super_admin", None
    await tracker.session.commit()
    user = UserRepository._to_entity(tracker.actor)
    with pytest.raises(TravelTrackerError) as error:
        await TravelTrackerService(tracker.session, user).mark(
            tracker.group.id,
            TrackerMarkRequest(track="visa", marked=True, passenger_ids=[tracker.passengers[0].id]),
        )
    assert error.value.status_code == 404


async def test_preview_exact_unique_name_passport_priority_and_no_mutation(tracker):
    content = b"Name,Passport Number\n asha  rao ,\nSame Name,\nWrong name,P1\nNobody,\nAsha Rao,\n"
    preview = await tracker.service.preview(tracker.group.id, content=content, filename="names.csv")
    assert [row.status for row in preview.rows] == [
        "matched",
        "ambiguous",
        "matched",
        "unmatched",
        "duplicate",
    ]
    assert preview.passenger_ids == [tracker.passengers[0].id, tracker.passengers[1].id]
    assert not (await tracker.session.scalars(select(TravelTrackerModel))).all()
    result = await tracker.service.mark(
        tracker.group.id,
        TrackerMarkRequest(
            track=preview.track, marked=preview.marked, passenger_ids=preview.passenger_ids
        ),
    )
    assert result.updated_count == 2


async def test_export_full_details_formula_safe_and_roundtrip(tracker):
    passenger = tracker.passengers[0]
    passenger.client_name = '=HYPERLINK("https://example.test")'
    passenger.staff_metadata = {
        "zone_name": "West",
        "designation": "Engineer",
        "malicious": "\t=1+1",
    }
    await tracker.session.commit()
    await tracker.service.mark(
        tracker.group.id,
        TrackerMarkRequest(track="visa", marked=True, passenger_ids=[passenger.id]),
    )
    content, filename = await tracker.service.export(tracker.group.id, status="marked")
    assert filename.endswith("-visa-marked.xlsx")
    workbook = load_workbook(io.BytesIO(content), data_only=False)
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    values = dict(zip(headers, [cell.value for cell in sheet[2]], strict=True))
    assert (
        sheet.max_row == 2 and values["Visa Applied"] is True and values["Flight Booked"] is False
    )
    assert values["Passport Number"] == "P0" and values["Nationality"] == "IND"
    assert values["Staff Code"] == "E42" and values["Imported: designation"] == "Engineer"
    assert values["Question: Meal"] == "Veg" and values["Detail: Employee ID"] == "EMP42"
    assert values["Full Name"].startswith("'=") and values["Imported: malicious"].startswith("'\t=")
    assert not any(cell.data_type == "f" for row in sheet for cell in row)
    workbook.close()
    preview = await tracker.service.preview(
        tracker.group.id, content=content, filename="export.xlsx", track="flight"
    )
    assert preview.matched_count == 1 and preview.passenger_ids == [passenger.id]
    pending, _ = await tracker.service.export(tracker.group.id, status="pending")
    book = load_workbook(io.BytesIO(pending))
    assert book.active.max_row == 6
    book.close()


async def test_composite_fk_rejects_cross_tenant_tracking(tracker):
    tracker.session.add(
        TravelTrackerModel(
            passenger_id=tracker.passengers[0].id,
            group_id=tracker.foreign.id,
            agency_id=tracker.foreign.agency_id,
            visa_applied=True,
        )
    )
    with pytest.raises(IntegrityError):
        await tracker.session.flush()
    await tracker.session.rollback()


async def test_mcp_actor_fence_does_not_upgrade_owner_lock(tracker, monkeypatch):
    from unittest.mock import AsyncMock

    from app.infrastructure.travel_tracker import service as service_module

    upgrade = AsyncMock(side_effect=AssertionError("MCP must retain its existing SHARE fence"))
    monkeypatch.setattr(service_module, "lock_tracker_actor", upgrade)
    service = TravelTrackerService(tracker.session, tracker.user, actor_already_fenced=True)
    result = await service.mark(
        tracker.group.id,
        TrackerMarkRequest(track="visa", marked=True, passenger_ids=[tracker.passengers[0].id]),
    )
    assert result.updated_count == 1
    upgrade.assert_not_awaited()


async def test_managed_mcp_transaction_rolls_back_marks_and_audit_together(tracker):
    result = await tracker.service.mark(
        tracker.group.id,
        TrackerMarkRequest(track="visa", marked=True, passenger_ids=[tracker.passengers[0].id]),
        commit=False,
    )
    assert result.updated_count == 1 and tracker.session.in_transaction()
    assert len((await tracker.session.scalars(select(TravelTrackerModel))).all()) == 1
    assert len((await tracker.session.scalars(select(AuditLogModel))).all()) == 1
    await tracker.session.rollback()
    assert not (await tracker.session.scalars(select(TravelTrackerModel))).all()
    assert not (await tracker.session.scalars(select(AuditLogModel))).all()


async def test_http_read_write_preview_export_and_response_contract(tracker, client):
    from app.presentation.dependencies.auth import get_current_active_user

    app = client._transport.app
    app.dependency_overrides[get_current_active_user] = lambda: tracker.user
    base = f"/api/v1/travel-tracker/groups/{tracker.group.id}"
    result = await client.get(base, params={"track": "flight", "status": "pending", "page_size": 2})
    assert result.status_code == 200 and "no-store" in result.headers["cache-control"]
    assert result.json()["total"] == 6 and len(result.json()["passengers"]) == 2
    result = await client.post(
        base + "/import/preview",
        data={"track": "flight", "marked": "true"},
        files={"file": ("names.csv", b"Name\nAsha Rao\n", "text/csv")},
    )
    assert result.status_code == 200 and result.json()["matched_count"] == 1
    result = await client.patch(
        base + "/marks",
        json={"track": "flight", "marked": True, "passenger_ids": result.json()["passenger_ids"]},
    )
    assert result.status_code == 200 and result.json()["updated_count"] == 1
    export = await client.get(base + "/export", params={"track": "flight", "status": "marked"})
    assert export.status_code == 200 and export.content.startswith(b"PK")
    assert "spreadsheetml.sheet" in export.headers["content-type"]
    invalid = await client.patch(
        base + "/marks",
        json={
            "track": "flight",
            "marked": "false",
            "passenger_ids": [str(tracker.passengers[0].id)],
        },
    )
    assert invalid.status_code == 422
    responses = app.openapi()["paths"]["/api/v1/travel-tracker/groups/{group_id}/export"]["get"][
        "responses"
    ]["200"]["content"]
    assert "application/json" not in responses


@pytest.mark.parametrize(
    "body",
    [
        {"track": "visa", "marked": "false", "passenger_ids": [str(uuid.uuid4())]},
        {"track": "visa", "marked": True},
        {"track": "visa", "marked": True, "selection": {"status": "pending"}},
        {
            "track": "visa",
            "marked": True,
            "selection": {"status": "pending"},
            "expected_count": 1,
            "passenger_ids": [str(uuid.uuid4())],
        },
        {"track": "visa", "marked": True, "passenger_ids": [str(uuid.uuid4())] * 1001},
    ],
)
def test_mark_contract_rejects_unbounded_or_ambiguous_requests(body):
    with pytest.raises(ValidationError):
        TrackerMarkRequest.model_validate(body)
