from __future__ import annotations

import asyncio
import hashlib
import json
import uuid
from io import BytesIO
from threading import Event
from types import SimpleNamespace
from unittest.mock import AsyncMock

import pytest
from fastapi import HTTPException, UploadFile
from openpyxl import load_workbook
from PIL import Image
from sqlalchemy import select, update
from starlette.datastructures import Headers

from app.domain.entities.entities import User, UserRole
from app.domain.exceptions.exceptions import AuthorizationError, ImageValidationError
from app.infrastructure.database.ecr_models import EcrBatchModel, EcrItemModel
from app.infrastructure.database.models import AgencyModel, UserModel, UserSecurityStateModel
from app.infrastructure.export.ecr_excel_exporter import build_ecr_workbook
from app.presentation.api.v1.routes import ecr_checker as routes
from app.presentation.api.v1.schemas.ecr_schemas import CreateEcrBatch


def _user(role=UserRole.AGENCY_ADMIN, agency_id=None):
    user_id = uuid.uuid4()
    return User(
        id=user_id,
        email=f"{user_id}@test.invalid",
        hashed_password="test-hash",
        full_name="ECR test user",
        agency_id=agency_id or uuid.uuid4(),
        role=role,
    )


async def _seed_user(session, user):
    if await session.get(AgencyModel, user.agency_id) is None:
        session.add(
            AgencyModel(
                id=user.agency_id, name="ECR agency", email=f"{user.agency_id}@test.invalid"
            )
        )
    if await session.get(UserModel, user.id) is None:
        session.add(
            UserModel(
                id=user.id,
                email=user.email,
                full_name=user.full_name,
                hashed_password=user.hashed_password,
                agency_id=user.actual_agency_id if user.actual_role is not None else user.agency_id,
                role=(user.actual_role or user.role).value,
                is_active=True,
            )
        )
        session.add(UserSecurityStateModel(user_id=user.id, session_version=user.session_version))
    await session.commit()


def _image():
    data = BytesIO()
    Image.new("RGB", (400, 300), "white").save(data, format="JPEG")
    return data.getvalue()


def _file(name="scan.jpg", content=None):
    return UploadFile(
        filename=name,
        file=BytesIO(content if content is not None else _image()),
        headers=Headers({"content-type": "image/jpeg"}),
    )


@pytest.fixture
def storage(monkeypatch):
    storage = SimpleNamespace(upload_file=AsyncMock(), delete_files=AsyncMock())
    security = SimpleNamespace(
        validate_image=AsyncMock(
            return_value=SimpleNamespace(content=_image(), content_type="image/jpeg", format="JPEG")
        )
    )
    monkeypatch.setattr(routes, "MinioStorageRepository", lambda: storage)
    monkeypatch.setattr(routes, "UploadSecurityService", lambda: security)
    monkeypatch.setattr(routes, "_dispatch", AsyncMock())
    return storage, security


async def _create(session, user, count=1):
    await _seed_user(session, user)
    return await routes.create_ecr_batch(
        CreateEcrBatch(title="Tomorrow", expected_count=count), user, session
    )


async def test_upload_preserves_duplicate_filenames_and_replays_idempotently(db_session, storage):
    user = _user()
    batch = await _create(db_session, user, 2)
    ids = [str(uuid.uuid4()), str(uuid.uuid4())]
    result = await routes.upload_ecr_items(
        batch.batch_id, user, db_session, [_file(), _file()], json.dumps(ids)
    )
    assert result.total_count == 2
    assert all(
        call.kwargs["max_dimension"] == routes.get_settings().ecr_image_max_dimension
        for call in storage[1].validate_image.await_args_list
    )
    assert [item.original_filename for item in result.items] == ["scan.jpg", "scan.jpg"]
    assert len({item.id for item in result.items}) == 2
    replay = await routes.upload_ecr_items(
        batch.batch_id, user, db_session, [_file(), _file()], json.dumps(ids)
    )
    assert replay.total_count == 2
    assert storage[0].upload_file.await_count == 2
    with pytest.raises(HTTPException) as error:
        await routes.upload_ecr_items(
            batch.batch_id, user, db_session, [_file(content=b"changed")], json.dumps(ids[:1])
        )
    assert error.value.status_code == 409


async def test_start_requires_every_file_then_dispatches_only_once(db_session, storage):
    user = _user()
    batch = await _create(db_session, user, 2)
    with pytest.raises(HTTPException) as error:
        await routes.start_ecr_batch(batch.batch_id, user, db_session)
    assert error.value.status_code == 409
    await routes.upload_ecr_items(
        batch.batch_id,
        user,
        db_session,
        [_file(), _file()],
        json.dumps([str(uuid.uuid4()), str(uuid.uuid4())]),
    )
    first = await routes.start_ecr_batch(batch.batch_id, user, db_session)
    second = await routes.start_ecr_batch(batch.batch_id, user, db_session)
    assert first.status == second.status == "queued"
    routes._dispatch.assert_awaited_once_with(batch.batch_id)


async def test_tenant_and_staff_access_are_enforced_for_read_and_export(db_session):
    owner = _user(UserRole.AGENCY_STAFF)
    batch = await _create(db_session, owner)
    for other in [_user(), _user(UserRole.AGENCY_STAFF, owner.agency_id)]:
        for action in [routes.get_ecr_batch, routes.export_ecr_batch]:
            with pytest.raises(HTTPException) as error:
                await action(batch.batch_id, other, db_session)
            assert error.value.status_code == 404
    manager = _user(UserRole.AGENCY_MANAGER, owner.agency_id)
    assert (
        await routes.get_ecr_batch(batch.batch_id, manager, db_session)
    ).batch_id == batch.batch_id
    assert await routes.list_ecr_batches(_user(), db_session) == []


@pytest.mark.parametrize("role", [UserRole.AGENCY_COORDINATOR, UserRole.CLIENT_MANAGER])
async def test_unprivileged_roles_cannot_create_batches(db_session, role):
    with pytest.raises(HTTPException) as error:
        await _create(db_session, _user(role))
    assert error.value.status_code == 403


async def test_invalid_upload_keeps_error_row_without_sending_or_storing(db_session, storage):
    storage[1].validate_image.side_effect = ImageValidationError("not readable")
    user = _user()
    batch = await _create(db_session, user)
    response = await routes.upload_ecr_items(
        batch.batch_id, user, db_session, [_file(content=b"bad")], json.dumps([str(uuid.uuid4())])
    )
    assert response.failed_count == 1
    assert response.na_count == 0
    assert response.items[0].reason == "invalid_image_upload"
    storage[0].upload_file.assert_not_awaited()


async def test_excel_waits_for_complete_batch_and_never_maps_failure_to_na(db_session, storage):
    user = _user()
    batch = await _create(db_session, user)
    with pytest.raises(HTTPException) as error:
        await routes.export_ecr_batch(batch.batch_id, user, db_session)
    assert error.value.status_code == 409
    await routes.upload_ecr_items(
        batch.batch_id, user, db_session, [_file()], json.dumps([str(uuid.uuid4())])
    )
    saved = await db_session.get(EcrBatchModel, batch.batch_id)
    saved.status = "completed_with_errors"
    item = (await db_session.execute(select(EcrItemModel))).scalar_one()
    item.status, item.reason = "failed", "timeout"
    await db_session.commit()
    response = await routes.export_ecr_batch(batch.batch_id, user, db_session)
    workbook = load_workbook(BytesIO(response.body))
    assert workbook.active["B2"].value == "ERROR"
    assert response.headers["cache-control"] == "no-store"


async def test_retry_only_failed_retained_images_preserves_completed_results(db_session, storage):
    user = _user()
    batch = await _create(db_session, user, 2)
    await routes.upload_ecr_items(
        batch.batch_id,
        user,
        db_session,
        [_file(), _file()],
        json.dumps([str(uuid.uuid4()), str(uuid.uuid4())]),
    )
    saved = await db_session.get(EcrBatchModel, batch.batch_id)
    saved.status = "completed_with_errors"
    items = list(
        (await db_session.execute(select(EcrItemModel).order_by(EcrItemModel.created_at))).scalars()
    )
    items[0].status, items[0].result = "completed", "ECR"
    items[1].status, items[1].reason = "failed", "timeout"
    await db_session.commit()
    response = await routes.retry_ecr_batch(batch.batch_id, user, db_session)
    assert response.status == "queued"
    assert items[0].status == "completed" and items[0].result == "ECR"
    assert items[1].status == "queued" and items[1].reason is None


def test_workbook_has_exact_two_columns_colors_duplicate_rows_and_literal_filenames():
    data = build_ecr_workbook(
        [
            ("duplicate.jpg", "ECR"),
            ("duplicate.jpg", "NA"),
            ('=HYPERLINK("https://example.invalid")', "REVIEW"),
            ("error.jpg", "ERROR"),
        ]
    )
    workbook = load_workbook(BytesIO(data))
    sheet = workbook.active
    assert sheet.max_column == 2 and sheet.max_row == 5
    assert [cell.value for cell in sheet[1]] == ["File name", "ECR"]
    assert sheet["B2"].font.color.rgb == "00FF0000"
    assert sheet["B3"].font.color.rgb == "00000000"
    assert sheet["A2"].value == sheet["A3"].value == "duplicate.jpg"
    assert sheet["A4"].data_type == "s"
    assert sheet["A4"].value.startswith("=HYPERLINK")
    assert sheet["B5"].value == "ERROR"


def test_client_ids_reject_duplicates_malformed_and_count_mismatches():
    value = str(uuid.uuid4())
    for raw in ["{}", '"abc"', "[null]", "[]", json.dumps([value, value])]:
        with pytest.raises(HTTPException):
            routes._parse_client_ids(raw, 1 if raw != json.dumps([value, value]) else 2)


async def test_upload_releases_database_transaction_before_scanning_and_storage(
    db_session, storage
):
    user = _user()
    batch = await _create(db_session, user)

    async def check_released(*args, **kwargs):
        assert not db_session.in_transaction()

    storage[0].upload_file.side_effect = check_released
    storage[1].validate_image.side_effect = check_released
    response = await routes.upload_ecr_items(
        batch.batch_id, user, db_session, [_file()], json.dumps([str(uuid.uuid4())])
    )
    assert response.total_count == 1
    storage[0].upload_file.assert_awaited_once()
    storage[1].validate_image.assert_awaited_once()


async def test_late_replay_race_preserves_winning_row_and_object(db_session, storage):
    user = _user()
    batch = await _create(db_session, user)
    client_id = uuid.uuid4()
    winner_id = uuid.uuid4()
    winner_key = f"ecr-checks/{user.agency_id}/{batch.batch_id}/{winner_id}.jpg"

    async def concurrent_commit(content, key, content_type):
        assert not db_session.in_transaction()
        assert key != winner_key
        db_session.add(
            EcrItemModel(
                id=winner_id,
                batch_id=batch.batch_id,
                client_id=client_id,
                original_filename="scan.jpg",
                content_type="image/jpeg",
                object_key=winner_key,
                sha256=hashlib.sha256(_image()).hexdigest(),
            )
        )
        await db_session.commit()

    storage[0].upload_file.side_effect = concurrent_commit
    response = await routes.upload_ecr_items(
        batch.batch_id, user, db_session, [_file()], json.dumps([str(client_id)])
    )
    assert response.total_count == 1
    assert response.items[0].id == winner_id
    assert (await db_session.get(EcrItemModel, winner_id)).object_key == winner_key
    # The loser is an attempt-unique orphan. Scheduled reconciliation proves
    # it unreferenced before deleting it, even after an uncertain commit.
    storage[0].delete_files.assert_not_awaited()


async def test_late_capacity_race_rejects_only_losing_attempt(db_session, storage):
    user = _user()
    batch = await _create(db_session, user)
    winner_id = uuid.uuid4()
    winner_key = f"ecr-checks/{user.agency_id}/{batch.batch_id}/{winner_id}.jpg"
    staged_keys = []

    async def concurrent_commit(content, key, content_type):
        staged_keys.append(key)
        db_session.add(
            EcrItemModel(
                id=winner_id,
                batch_id=batch.batch_id,
                client_id=uuid.uuid4(),
                original_filename="different.jpg",
                content_type="image/jpeg",
                object_key=winner_key,
                sha256=hashlib.sha256(_image()).hexdigest(),
            )
        )
        await db_session.commit()

    storage[0].upload_file.side_effect = concurrent_commit
    with pytest.raises(HTTPException) as error:
        await routes.upload_ecr_items(
            batch.batch_id, user, db_session, [_file()], json.dumps([str(uuid.uuid4())])
        )
    assert error.value.status_code == 409
    rows = list((await db_session.scalars(select(EcrItemModel))).all())
    assert len(rows) == 1 and rows[0].id == winner_id
    storage[0].delete_files.assert_awaited_once_with(staged_keys)
    assert winner_key not in staged_keys


async def test_lost_commit_acknowledgement_never_deletes_a_committed_image(
    db_session, storage, monkeypatch
):
    user = _user()
    batch = await _create(db_session, user)
    real_commit = db_session.commit

    async def lose_final_ack():
        committing_image = any(isinstance(row, EcrItemModel) for row in db_session.new)
        await real_commit()
        if committing_image:
            raise RuntimeError("simulated lost database commit acknowledgement")

    monkeypatch.setattr(db_session, "commit", lose_final_ack)
    with pytest.raises(RuntimeError, match="lost database commit acknowledgement"):
        await routes.upload_ecr_items(
            batch.batch_id, user, db_session, [_file()], json.dumps([str(uuid.uuid4())])
        )
    rows = list((await db_session.scalars(select(EcrItemModel))).all())
    assert len(rows) == 1
    assert rows[0].object_key == storage[0].upload_file.await_args.args[1]
    storage[0].delete_files.assert_not_awaited()


async def test_disabled_agency_cannot_read_mutate_or_export_ecr(db_session, storage):
    user = _user()
    batch = await _create(db_session, user)
    await db_session.execute(
        update(AgencyModel).where(AgencyModel.id == user.agency_id).values(is_active=False)
    )
    await db_session.commit()
    assert await routes.list_ecr_batches(user, db_session) == []
    for action in [routes.get_ecr_batch, routes.export_ecr_batch]:
        with pytest.raises(HTTPException) as error:
            await action(batch.batch_id, user, db_session)
        assert error.value.status_code == 404
    for action in [routes.start_ecr_batch, routes.retry_ecr_batch]:
        with pytest.raises(HTTPException) as error:
            await action(batch.batch_id, user, db_session)
        assert error.value.status_code == 403
    with pytest.raises(HTTPException) as error:
        await routes.create_ecr_batch(CreateEcrBatch(expected_count=1), user, db_session)
    assert error.value.status_code == 403
    with pytest.raises(HTTPException) as error:
        await routes.upload_ecr_items(
            batch.batch_id, user, db_session, [_file()], json.dumps([str(uuid.uuid4())])
        )
    assert error.value.status_code == 404
    storage[0].upload_file.assert_not_awaited()


async def test_role_revoked_during_image_staging_is_rechecked_before_commit(db_session, storage):
    user = _user()
    batch = await _create(db_session, user)

    async def revoke_role(content, key, content_type):
        await db_session.execute(
            update(UserModel)
            .where(UserModel.id == user.id)
            .values(role=UserRole.AGENCY_COORDINATOR.value)
        )
        await db_session.commit()

    storage[0].upload_file.side_effect = revoke_role
    with pytest.raises(HTTPException) as error:
        await routes.upload_ecr_items(
            batch.batch_id, user, db_session, [_file()], json.dumps([str(uuid.uuid4())])
        )
    assert error.value.status_code == 403
    assert list((await db_session.scalars(select(EcrItemModel))).all()) == []
    storage[0].delete_files.assert_awaited_once()


async def test_restricted_super_admin_keeps_selected_agency_and_staff_ownership(db_session):
    user = _user(UserRole.AGENCY_STAFF)
    user.actual_role = UserRole.SUPER_ADMIN
    user.actual_agency_id = None
    batch = await _create(db_session, user)
    stored = await db_session.get(EcrBatchModel, batch.batch_id)
    assert stored.agency_id == user.agency_id
    assert stored.created_by_user_id == user.id
    assert (
        await routes.list_ecr_batches(_user(UserRole.AGENCY_STAFF, user.agency_id), db_session)
        == []
    )


async def test_session_revoked_during_image_staging_is_rechecked_before_commit(db_session, storage):
    user = _user()
    batch = await _create(db_session, user)

    async def revoke_session(content, key, content_type):
        await db_session.execute(
            update(UserSecurityStateModel)
            .where(UserSecurityStateModel.user_id == user.id)
            .values(session_version=2)
        )
        await db_session.commit()

    storage[0].upload_file.side_effect = revoke_session
    with pytest.raises(AuthorizationError):
        await routes.upload_ecr_items(
            batch.batch_id, user, db_session, [_file()], json.dumps([str(uuid.uuid4())])
        )
    assert list((await db_session.scalars(select(EcrItemModel))).all()) == []
    storage[0].delete_files.assert_awaited_once()


async def test_history_summary_uses_aggregated_outcomes_and_keeps_empty_batches(db_session):
    user = _user()
    batch = await _create(db_session, user, 5)
    empty = await _create(db_session, user)
    for status, result in [
        ("completed", "ECR"),
        ("completed", "NA"),
        ("completed", "NEEDS_REVIEW"),
        ("failed", None),
        ("queued", None),
    ]:
        db_session.add(
            EcrItemModel(
                batch_id=batch.batch_id,
                client_id=uuid.uuid4(),
                original_filename="scan.jpg",
                content_type="image/jpeg",
                sha256="0" * 64,
                status=status,
                result=result,
            )
        )
    await db_session.commit()
    summaries = {item.batch_id: item for item in await routes.list_ecr_batches(user, db_session)}
    filled = summaries[batch.batch_id]
    assert (
        filled.total_count,
        filled.processed_count,
        filled.ecr_count,
        filled.na_count,
        filled.review_count,
        filled.failed_count,
    ) == (5, 4, 1, 1, 1, 1)
    assert summaries[empty.batch_id].total_count == 0
    assert summaries[empty.batch_id].processed_count == 0


async def test_upload_admission_limits_concurrent_operations_to_one():
    active = peak = 0

    async def operation():
        nonlocal active, peak
        active += 1
        peak = max(peak, active)
        await asyncio.sleep(0.001)
        active -= 1
        return "done"

    results = await asyncio.wait_for(
        asyncio.gather(*(routes._run_admitted_upload(operation) for _ in range(12))), timeout=3
    )
    assert results == ["done"] * 12
    assert peak == 1 and active == 0


async def test_cancelled_upload_retains_admission_until_decoder_thread_settles():
    decoder_release = Event()
    decoder_started = asyncio.Event()
    second_started = asyncio.Event()
    loop = asyncio.get_running_loop()

    def decode():
        loop.call_soon_threadsafe(decoder_started.set)
        assert decoder_release.wait(timeout=5)

    async def first_operation():
        await asyncio.to_thread(decode)
        return "first"

    async def second_operation():
        second_started.set()
        return "second"

    first = asyncio.create_task(routes._run_admitted_upload(first_operation))
    second = None
    try:
        await asyncio.wait_for(decoder_started.wait(), timeout=2)
        second = asyncio.create_task(routes._run_admitted_upload(second_operation))
        first.cancel()
        await asyncio.sleep(0.01)
        first.cancel()  # Repeated disconnect/shutdown cancellation cannot free the slot either.
        await asyncio.sleep(0.01)
        assert not first.done()
        assert not second_started.is_set()
        decoder_release.set()
        with pytest.raises(asyncio.CancelledError):
            await asyncio.wait_for(first, timeout=3)
        assert await asyncio.wait_for(second, timeout=3) == "second"
    finally:
        decoder_release.set()
        await asyncio.gather(
            first, *([second] if second is not None else []), return_exceptions=True
        )


async def test_cancelled_waiting_upload_does_not_start_work():
    release = asyncio.Event()
    admitted = asyncio.Event()

    async def first_operation():
        admitted.set()
        await release.wait()
        return "first"

    queued_operation = AsyncMock()
    first = asyncio.create_task(routes._run_admitted_upload(first_operation))
    await asyncio.wait_for(admitted.wait(), timeout=2)
    queued = asyncio.create_task(routes._run_admitted_upload(queued_operation))
    await asyncio.sleep(0)
    queued.cancel()
    try:
        with pytest.raises(asyncio.CancelledError):
            await asyncio.wait_for(queued, timeout=2)
        queued_operation.assert_not_awaited()
    finally:
        release.set()
        await asyncio.wait_for(first, timeout=2)


async def test_waiting_upload_body_is_not_read_until_admission(monkeypatch):
    release = asyncio.Event()
    admitted = asyncio.Event()
    first_file = SimpleNamespace(read=AsyncMock(return_value=b"first"))
    waiting_file = SimpleNamespace(read=AsyncMock(return_value=b"second"))

    async def admitted_upload(*, files, **kwargs):
        content = await files[0].read()
        if content == b"first":
            admitted.set()
            await release.wait()
        return content

    monkeypatch.setattr(routes, "_upload_ecr_items_admitted", admitted_upload)
    session = SimpleNamespace(commit=AsyncMock())
    first = asyncio.create_task(
        routes.upload_ecr_items(uuid.uuid4(), _user(), session, [first_file], "[]")
    )
    waiting = None
    try:
        await asyncio.wait_for(admitted.wait(), timeout=2)
        waiting = asyncio.create_task(
            routes.upload_ecr_items(uuid.uuid4(), _user(), session, [waiting_file], "[]")
        )
        await asyncio.sleep(0.01)
        waiting_file.read.assert_not_awaited()
        assert (
            session.commit.await_count == 2
        )  # Neither request holds a DB transaction while waiting.
        release.set()
        assert await asyncio.wait_for(first, timeout=2) == b"first"
        assert await asyncio.wait_for(waiting, timeout=2) == b"second"
        waiting_file.read.assert_awaited_once_with()
    finally:
        release.set()
        await asyncio.wait_for(
            asyncio.gather(
                first, *([waiting] if waiting is not None else []), return_exceptions=True
            ),
            timeout=3,
        )
