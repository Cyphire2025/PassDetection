"""PostgreSQL source admission at the minimum deployed profile, not capacity.

Keep the UUID-owned schema and its 5,000/100 synthetic cohorts after execution.
Schema comes from ORM metadata: this is not migration, runtime-role, HTTP, object
storage, Linux memory or combined-load qualification. Existing export race tests
separately cover a busy contender recovering the same durable operation.
"""

from __future__ import annotations

import io
import os
import uuid
from collections import Counter
from dataclasses import replace
from types import SimpleNamespace

import pytest
import pytest_asyncio
from openpyxl import load_workbook
from sqlalchemy import URL, Text, cast, event, func, insert, select, text, update
from sqlalchemy.ext.asyncio import async_sessionmaker, create_async_engine
from sqlalchemy.pool import NullPool
from sqlalchemy.schema import CreateSchema

from app.application.mcp.artifacts import ArtifactError
from app.application.mcp.authorization import MCPAuthorizationService
from app.application.mcp.exports import ExcelExportRequest, MCPExcelExportService
from app.core.config.mcp import MCPSettings
from app.infrastructure.database.mcp_models import MCPControlModel
from app.infrastructure.database.model_registry import Base
from app.infrastructure.database.models import (
    AgencyModel,
    ClientGroupModel,
    PassportSubmissionModel,
)
from app.infrastructure.repositories.user_repository import UserRepository
from app.presentation.api.v1.routes.passport_routes.excel_exports import export_passports_by_group
from app.presentation.api.v1.routes.passport_routes.selected_exports import (
    export_selected_passports,
)
from app.presentation.api.v1.schemas.passport_schemas import ExportSelectedPassportsRequest
from app.presentation.mcp.export_tools import excel_support
from tests.integration.test_mcp_operations import seed_identity

ROWS, BYTES = 100, 1024 * 1024
pytestmark = [
    pytest.mark.service_integration,
    pytest.mark.skipif(os.getenv("RUN_SERVICE_INTEGRATION") != "1", reason="isolated PostgreSQL required"),
]


@pytest_asyncio.fixture(scope="module", loop_scope="module")
async def retained_cohorts(test_settings, record_testsuite_property):
    database = os.environ["POSTGRES_DB"]
    host = os.environ.get("POSTGRES_HOST", "127.0.0.1")
    if host not in {"localhost", "127.0.0.1"} or not database.startswith("passdetection_ci_"):
        pytest.fail("Source admission proof requires a dedicated loopback CI database")
    schema = "mcp_source_" + uuid.uuid4().hex
    record_testsuite_property("retained_postgresql_component_schema", schema)
    engine = create_async_engine(
        URL.create("postgresql+asyncpg", username=os.environ["POSTGRES_USER"],
                   password=os.environ.get("POSTGRES_PASSWORD"), host=host,
                   port=int(os.environ.get("POSTGRES_PORT", "5432")), database=database),
        poolclass=NullPool,
        # Existing pg_trgm lives in public; all ORM tables remain explicitly
        # schema-translated, including foreign keys and source-budget queries.
        connect_args={"server_settings": {"search_path": f"{schema},public", "statement_timeout": "30000", "lock_timeout": "2000"}},
        execution_options={"schema_translate_map": {None: schema}},
    )
    sessions = async_sessionmaker(engine, expire_on_commit=False)
    settings = test_settings.model_copy(update={"mcp": MCPSettings(
        enabled=True, enabled_capabilities=["mcp:read", "mcp:export"],
        export_families=["passport_excel"], export_source_row_limit=ROWS,
        export_source_byte_limit=BYTES,
    )})
    try:
        async with engine.begin() as connection:
            await connection.execute(CreateSchema(schema))
            assert await connection.scalar(text("SELECT current_schema()")) == schema
            await connection.run_sync(Base.metadata.create_all)
        async with sessions() as session:
            session.add(MCPControlModel(id=1, enabled=True))
            user, grants, tokens = await seed_identity(session, settings)
            for grant in grants:
                grant.capabilities = ["mcp:read", "mcp:export"]
            agency = AgencyModel(id=uuid.uuid4(), name="Source admission fixture", email=f"{schema}@example.test")
            session.add(agency)
            await session.flush()
            groups, identifiers = {}, {}
            for cohort, count in (("large", 5000), ("small", 100)):
                group = ClientGroupModel(id=uuid.uuid4(), agency_id=agency.id,
                    name=f"Synthetic {cohort}", token=uuid.uuid4().hex)
                session.add(group)
                await session.flush()
                groups[cohort] = group.id
                identifiers[cohort] = [uuid.uuid4() for _ in range(count)]
                await session.execute(insert(PassportSubmissionModel), [
                    {"id": identifier, "agency_id": agency.id, "group_id": group.id,
                     "client_name": f"Synthetic {index}", "image_s3_key": "private/never-read",
                     "status": "staff_approved", "confirmed_fields": {
                         "given_names": f"Synthetic {index}", "surname": cohort,
                         "passport_number": "DUPLICATE" if index < 50 else f"TEST{index}",
                     }} for index, identifier in enumerate(identifiers[cohort])
                ])
            await session.commit()
        yield SimpleNamespace(sessions=sessions, settings=settings, schema=schema,
                              user_id=user.id, token=tokens[0], agency_id=agency.id,
                              groups=groups, identifiers=identifiers)
    finally:
        # Deliberately retain the schema and committed synthetic records.
        await engine.dispose()


def request_for(fixture, cohort, *, count=ROWS):
    return ExcelExportRequest(
        agency_id=fixture.agency_id, group_ids=[fixture.groups[cohort]],
        selection="selected_passports" if cohort == "large" else "group",
        submission_ids=fixture.identifiers[cohort][:count] if cohort == "large" else [],
    )


async def service_and_principal(fixture, session):
    principal = await MCPAuthorizationService(session, fixture.settings).verify_access(fixture.token)
    return MCPExcelExportService(session, fixture.settings, excel_support()), principal


def workbook_rows(content):
    workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False)
    try:
        assert workbook.sheetnames == ["Passport Submissions"]
        sheet = workbook["Passport Submissions"]
        header = next(sheet.iter_rows(min_row=4, max_row=4, values_only=True))
        rows = Counter(tuple(row) for row in sheet.iter_rows(min_row=5, values_only=True)
                       if any(value is not None for value in row))
        return header, rows
    finally:
        workbook.close()


@pytest.mark.parametrize("cohort", ["large", "small"])
async def test_exact_100_selection_has_canonical_workbook_parity_and_bounded_orm(retained_cohorts, cohort):
    f = retained_cohorts
    selected = set(f.identifiers[cohort][:ROWS])
    loaded = set()

    def loaded_passport(row, _context):
        loaded.add(row.id)

    async with f.sessions() as session:
        service, principal = await service_and_principal(f, session)
        event.listen(PassportSubmissionModel, "load", loaded_passport)
        try:
            observed = await service.inspect(principal, request_for(f, cohort))
            prepared = await service.prepare(principal, request_for(f, cohort))
        finally:
            event.remove(PassportSubmissionModel, "load", loaded_passport)
        assert observed["passenger_count"] == ROWS and observed["pending_recipient_count"] == 0
        assert observed["maximum_source_rows_per_family"] == ROWS
        assert observed["maximum_source_bytes"] == BYTES
        assert observed["expected_revision"] == prepared.revision
        assert loaded == selected
        assert {row.id for row in prepared.submissions} == selected
        actual = workbook_rows(await prepared.render())
        actor = replace(await UserRepository(session).get_by_id(f.user_id), agency_id=f.agency_id)
        if cohort == "large":
            response = await export_selected_passports(
                ExportSelectedPassportsRequest(submission_ids=f.identifiers[cohort][:ROWS]),
                current_user=actor, session=session,
            )
        else:
            response = await export_passports_by_group(
                f.groups[cohort], export_mode="all", baseline_export_id=None,
                request_id=uuid.uuid4(), supplemental_fields=None, group_by_field=None,
                agency_match_field=None, current_user=actor, session=session,
            )
        web = workbook_rows(b"".join([part async for part in response.body_iterator]))
        assert actual == web
        header, rows = actual
        indexes = [header.index(name) for name in ("GIVEN NAME", "SURNAME", "Passport Number")]
        identities = Counter()
        for row, count in rows.items():
            identities[tuple(str(row[index]).casefold() for index in indexes)] += count
        expected = Counter((f"synthetic {index}", cohort, "duplicate" if index < 50 else f"test{index}")
                           for index in range(ROWS))
        assert identities == expected and sum(rows.values()) == ROWS
        counts = dict((await session.execute(select(PassportSubmissionModel.group_id, func.count())
            .where(PassportSubmissionModel.group_id.in_(f.groups.values()))
            .group_by(PassportSubmissionModel.group_id))).all())
        assert counts == {f.groups["large"]: 5000, f.groups["small"]: 100}


@pytest.mark.parametrize("selection", ["101_selected", "5000_whole_group"])
async def test_excess_source_rows_rejected_before_passport_orm(retained_cohorts, selection):
    f = retained_cohorts
    request = request_for(f, "large", count=101) if selection == "101_selected" else ExcelExportRequest(
        agency_id=f.agency_id, group_ids=[f.groups["large"]],
    )

    def forbidden_load(*_args):
        raise AssertionError("Rejected source reached full passport ORM loading")

    async with f.sessions() as session:
        service, principal = await service_and_principal(f, session)
        event.listen(PassportSubmissionModel, "load", forbidden_load)
        try:
            with pytest.raises(ArtifactError, match="configured row limit") as denied:
                await service.inspect(principal, request)
            assert denied.value.status_code == 413
        finally:
            event.remove(PassportSubmissionModel, "load", forbidden_load)


async def test_cumulative_unicode_bytes_exceed_1mib_before_orm(retained_cohorts):
    f = retained_cohorts
    identifiers = f.identifiers["large"][:2]

    def forbidden_load(*_args):
        raise AssertionError("Oversized UTF-8 source reached full passport ORM loading")

    async with f.sessions() as session:
        # Individual values fit; their aggregate UTF-8 bytes exceed the exact
        # deployed ceiling while their aggregate character count remains below it.
        await session.execute(update(PassportSubmissionModel)
            .where(PassportSubmissionModel.id.in_(identifiers))
            .values(confirmed_fields={"oversized": "旅" * 180000}))
        value = cast(PassportSubmissionModel.confirmed_fields, Text)
        sizes = (await session.execute(select(func.max(func.octet_length(value)),
            func.sum(func.octet_length(value)), func.sum(func.char_length(value)))
            .where(PassportSubmissionModel.id.in_(identifiers)))).one()
        assert sizes[0] < BYTES < sizes[1] and sizes[2] < BYTES
        service, principal = await service_and_principal(f, session)
        event.listen(PassportSubmissionModel, "load", forbidden_load)
        try:
            with pytest.raises(ArtifactError, match="configured byte limit") as denied:
                await service.inspect(principal, request_for(f, "large"))
            assert denied.value.status_code == 413
        finally:
            event.remove(PassportSubmissionModel, "load", forbidden_load)
        # Roll back only this test's uncommitted oversized values. Retained seed
        # records, schemas, files and production resources are never deleted.
        await session.rollback()
