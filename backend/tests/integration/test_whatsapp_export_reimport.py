"""Real selected-passport workbook round trips; no provider calls or real PII."""

import uuid
from io import BytesIO

import pytest
from openpyxl import Workbook

from app.domain.entities.entities import PassportSubmission
from app.domain.services.traveller_identity import compare_identity, identity_evidence
from app.infrastructure.export.passport_excel_exporter import PassportExcelExporter
from app.presentation.api.v1.routes.whatsapp_contact_import import _parse_excel_contact_bytes
from app.presentation.api.v1.routes.whatsapp_contact_support import (
    _excel_header_label,
    _is_excel_phone_header,
)


def selected_passports(count: int, *, shared_phone: bool) -> list[PassportSubmission]:
    agency_id, group_id = uuid.uuid4(), uuid.uuid4()
    submissions = []
    for index in range(count):
        name = f"Traveller {index:03d} Synthetic"
        source = PassportSubmission.create(
            group_id=group_id,
            agency_id=agency_id,
            client_name=name,
            client_email=None,
            image_s3_key="synthetic/front.jpg",
        )
        source.submit_client_review(
            {
                "given_names": f"Traveller {index:03d}",
                "surname": "Synthetic",
                "passport_number": f"P{index:07d}",
                "date_of_birth": "1990-01-01",
                "nationality": "IND",
            },
            client_email=f"traveller{index}@example.test",
            client_phone=f"98{index // 2 if shared_phone else index:08d}",
            departure_city="Synthetic city",
        )
        submissions.append(source)
    return submissions


@pytest.mark.parametrize("shared_phone", [False, True])
def test_selected_passport_export_reimports_100_people_without_email_phone_candidates(shared_phone):
    submissions = selected_passports(100, shared_phone=shared_phone)
    content = PassportExcelExporter().export_group(
        submissions,
        group_name="Selected Passports",
        whatsapp_contacts={
            source.id: {"email": f"broadcast{index}@example.test", "phone": "9899999999"}
            for index, source in enumerate(submissions)
        },
    )
    parsed = _parse_excel_contact_bytes(content, filename="selected-passports.xlsx")
    assert len(parsed.contacts) == (50 if shared_phone else 100)
    assert parsed.rejected_counts == ({"duplicate_phone": 50} if shared_phone else {})
    assert parsed.rejected_count == (50 if shared_phone else 0)
    all_rows = [*parsed.contacts, *parsed.rejected_rows]
    assert len(all_rows) == 100
    by_passport = {source.confirmed_fields["passport_number"]: source for source in submissions}
    for row in all_rows:
        fields = row.imported_fields
        source = by_passport[fields["passport_number"]]
        index = int(fields["passport_number"][1:])
        assert fields["whatsapp_email"] == f"broadcast{index}@example.test"
        assert fields["client_email"] == source.client_email
        # The submitted email participates in identity matching; the separately
        # labelled broadcast email remains metadata, not contradictory identity.
        imported = identity_evidence(fields)
        assert imported.get("email") == frozenset({source.client_email})
        assert not imported.conflicts
        saved = identity_evidence(
            source.confirmed_fields,
            name=source.client_name,
            phone=source.client_phone,
            email=source.client_email,
        )
        assert compare_identity(imported, saved).matched


@pytest.mark.parametrize(
    "header",
    [
        "WhatsApp Email",
        "WhatsApp E-mail Address",
        "Upload Email",
        "Email",
        "Mobile Email",
        "Telephone e-mail",
        "Contact mail",
    ],
)
def test_email_headers_cannot_designate_a_delivery_phone(header):
    assert not _is_excel_phone_header(_excel_header_label(header))


@pytest.mark.parametrize(
    "header",
    [
        "Verified WhatsApp Numbers",
        "Emergency Mobile",
        "Office phone 2",
        "Primary telephone",
        "Personal WhatsApp contact",
        "Contact number",
    ],
)
def test_additional_phone_columns_remain_supported(header):
    assert _is_excel_phone_header(_excel_header_label(header))


def test_genuine_invalid_additional_phone_is_still_reported_and_email_metadata_retained():
    book, output = Workbook(), BytesIO()
    book.active.append(
        ["Name", "Primary WhatsApp Phone", "WhatsApp Email", "Upload Email", "Emergency Mobile"]
    )
    book.active.append(
        [
            "Synthetic Person",
            "9876543210",
            "broadcast@example.test",
            "traveller@example.test",
            "not-a-phone",
        ]
    )
    book.save(output)
    book.close()
    parsed = _parse_excel_contact_bytes(output.getvalue(), filename="contacts.xlsx")
    assert len(parsed.contacts) == 1
    assert parsed.rejected_counts == {"invalid_phone": 1}
    assert parsed.rejected_rows[0].raw_phone_number == "not-a-phone"
    assert parsed.contacts[0].imported_fields["whatsapp_email"] == "broadcast@example.test"
    assert parsed.contacts[0].imported_fields["client_email"] == "traveller@example.test"


@pytest.mark.parametrize("include_phone", [False, True])
@pytest.mark.parametrize("include_email_header", [False, True])
def test_phone_fallback_never_uses_numeric_email_local_part(include_phone, include_email_header):
    book, output = Workbook(), BytesIO()
    if include_email_header:
        book.active.append(["Name", "WhatsApp Email"])
    row = ["Synthetic Person", "9988776655@example.test"]
    if include_phone:
        row.append("9876543210")
    book.active.append(row)
    book.save(output)
    book.close()

    parsed = _parse_excel_contact_bytes(output.getvalue(), filename="unlabelled.xlsx")

    assert parsed.rejected_rows == []
    assert [contact.phone_number for contact in parsed.contacts] == (
        ["9876543210"] if include_phone else []
    )


def test_email_in_explicit_phone_column_remains_a_rejected_phone():
    book, output = Workbook(), BytesIO()
    book.active.append(["Name", "Phone"])
    book.active.append(["Synthetic Person", "9988776655@example.test"])
    book.save(output)
    book.close()

    parsed = _parse_excel_contact_bytes(output.getvalue(), filename="incorrect-phone.xlsx")

    assert parsed.contacts == []
    assert parsed.rejected_counts == {"invalid_phone": 1}
    assert parsed.rejected_rows[0].raw_phone_number == "9988776655@example.test"
