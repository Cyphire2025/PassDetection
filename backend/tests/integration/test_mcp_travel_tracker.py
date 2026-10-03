"""Canonical OAuth, device permissions, real tracker mutations and native Excel delivery."""

from __future__ import annotations

import hashlib
import io
import uuid
from contextlib import asynccontextmanager
from datetime import UTC, datetime

import pytest
from mcp.server.auth.provider import AccessToken
from mcp.server.mcpserver.exceptions import ToolError
from openpyxl import Workbook, load_workbook
from sqlalchemy import func, select

from app.application.mcp.artifacts import ArtifactError
from app.application.mcp.credentials import MCPAuthError
from app.application.mcp.operations import MCPOperationService
from app.application.mcp.travel_tracker import (
    MCPTravelTrackerExportService,
    TrackerExportRequest,
    tracker_export_operation,
)
from app.domain.mcp_read_sections import READ_TOOL_SECTIONS
from app.domain.mcp_section_permissions import WRITE_TOOL_SECTIONS
from app.infrastructure.database.mcp_models import MCPControlModel, MCPGrantModel
from app.infrastructure.database.mcp_operation_models import MCPOperationModel
from app.infrastructure.database.models import AuditLogModel
from app.infrastructure.database.travel_tracker_model import TravelTrackerModel
from app.presentation.mcp.server import ObservationalMCPServer, ReviewedMCPServer
from app.presentation.mcp.travel_tracker_tools import register_travel_tracker_tools
from tests.integration.test_mcp_artifacts import artifacts as artifacts
from tests.integration.test_mcp_exports import downloaded, person
from tests.integration.test_mcp_native_transfers import (
    native as native,
)
from tests.integration.test_mcp_native_transfers import (
    put,
    upload_ticket,
)


@pytest.fixture
async def tracker(native, monkeypatch):
    f = native
    control = await f.session.get(MCPControlModel, 1)
    control.allowed_read_sections = ["documents", "all_groups"]
    await f.session.commit()
    f.token, f.principal = await f.connect(["mcp:read", "mcp:change", "mcp:upload", "mcp:export"])
    f.client.headers["Authorization"] = f"Bearer {f.token}"
    f.passengers = [
        person(f, 1, status="pending_upload"),
        person(f, 2, status="staff_approved"),
        person(f, 3, status="pending_upload"),
    ]
    f.passengers[2].client_name = f.passengers[1].client_name
    f.session.add_all(f.passengers)
    await f.session.commit()
    f.app, f.server = f.client._transport.app, ReviewedMCPServer("Tracker fixture")

    @asynccontextmanager
    async def sessions():
        yield f.session

    f.app.state.mcp_session_factory, f.app.state.mcp_operations = sessions, {}
    register_travel_tracker_tools(f.server, f.app, f.settings)

    def current_token():
        return AccessToken(
            token=f.token,
            client_id=f.principal.client_id,
            scopes=list(f.principal.capabilities),
            subject=str(f.principal.user_id),
            resource=f.settings.mcp.resource,
            claims={"grant_id": str(f.principal.grant_id)},
        )

    monkeypatch.setattr("app.presentation.mcp.invocation.get_access_token", current_token)
    monkeypatch.setattr(
        "app.presentation.mcp.native_transfer_tools.get_access_token", current_token
    )

    async def call(name, **args):
        return (await f.server.call_tool(name, args)).structured_content

    f.call = call
    return f


async def test_real_roster_and_fast_independent_marks_have_one_receipt(tracker):
    f = tracker
    groups = await f.call("list_travel_tracker_groups")
    assert "groups" in groups, groups
    assert groups["groups"][0]["total"] == 3
    roster = await f.call("get_travel_tracker_roster", group_id=str(f.group.id), page_size=1)
    assert roster["next_page"] == 2 and roster["total"] == 3
    update = {
        "track": "visa",
        "marked": True,
        "passenger_ids": [str(f.passengers[0].id), str(f.passengers[1].id)],
    }
    args = {
        "group_id": str(f.group.id),
        "update": update,
        "idempotency_key": "tracker-real-mark-001",
    }
    first, repeated = (
        await f.call("set_travel_tracker_status", **args),
        await f.call("set_travel_tracker_status", **args),
    )
    assert first["receipt"] == repeated["receipt"]
    assert first["receipt"]["data"]["updated_count"] == 2
    marked = await f.call("get_travel_tracker_roster", group_id=str(f.group.id), status="marked")
    assert marked["total"] == 2 and all(
        row["visa_applied"] and not row["flight_booked"] for row in marked["passengers"]
    )
    corrected = await f.call(
        "set_travel_tracker_status",
        group_id=str(f.group.id),
        update={"track": "visa", "marked": False, "passenger_ids": [str(f.passengers[0].id)]},
        idempotency_key="tracker-correction-001",
    )
    assert corrected["receipt"]["data"]["updated_count"] == 1
    assert await f.session.scalar(select(func.count()).select_from(MCPOperationModel)) == 2
    changes = (
        await f.session.scalars(
            select(AuditLogModel).where(AuditLogModel.action == "travel_tracker.marks_changed")
        )
    ).all()
    assert len(changes) == 2


async def test_stale_bulk_count_blocks_all_changes_and_foreign_ids_are_atomic(tracker):
    f = tracker
    result = await f.call(
        "set_travel_tracker_status",
        group_id=str(f.group.id),
        update={
            "track": "flight",
            "marked": True,
            "selection": {"status": "pending"},
            "expected_count": 2,
        },
        idempotency_key="tracker-stale-count-001",
    )
    assert result["error"] == "tracker_selection_changed"
    await f.session.refresh(f.group)
    await f.session.refresh(f.passengers[0])
    foreign = await f.call(
        "set_travel_tracker_status",
        group_id=str(f.group.id),
        update={
            "track": "visa",
            "marked": True,
            "passenger_ids": [str(f.passengers[0].id), str(uuid.uuid4())],
        },
        idempotency_key="tracker-foreign-id-001",
    )
    assert "receipt" not in foreign
    assert await f.session.scalar(select(func.count()).select_from(TravelTrackerModel)) == 0


@pytest.mark.parametrize(
    "boundary",
    [
        "global_read",
        "device_read",
        "read_section",
        "global_write",
        "device_write",
        "write_section",
        "tool",
        "revoked",
    ],
)
async def test_live_access_boundaries_block_reads_or_writes(tracker, boundary):
    f = tracker
    control = await f.session.get(MCPControlModel, 1)
    grant = await f.session.get(MCPGrantModel, f.principal.grant_id)
    if boundary == "global_read":
        control.read_enabled = False
    elif boundary == "device_read":
        grant.read_enabled = False
    elif boundary == "read_section":
        grant.allowed_read_sections = ["documents"]
    elif boundary == "global_write":
        control.write_enabled = False
    elif boundary == "device_write":
        grant.write_enabled = False
    elif boundary == "write_section":
        grant.allowed_write_sections = ["documents"]
    elif boundary == "tool":
        control.allowed_write_tools = [
            name for name in control.allowed_write_tools if name != "set_travel_tracker_status"
        ]
    else:
        grant.revoked_at = datetime.now(UTC)
    await f.session.commit()
    if boundary in {"global_read", "device_read", "read_section"}:
        result = await f.call("get_travel_tracker_roster", group_id=str(f.group.id))
    else:
        result = await f.call(
            "set_travel_tracker_status",
            group_id=str(f.group.id),
            update={"track": "visa", "marked": True, "passenger_ids": [str(f.passengers[0].id)]},
            idempotency_key="tracker-denied-001",
        )
    assert result["error"] == "access_denied", result
    assert "passengers" not in result and "receipt" not in result
    assert await f.session.scalar(select(func.count()).select_from(TravelTrackerModel)) == 0


async def test_replay_revalidates_current_section_and_read_only_scope_cannot_write(tracker):
    f = tracker
    args = {
        "group_id": str(f.group.id),
        "update": {"track": "visa", "marked": True, "passenger_ids": [str(f.passengers[0].id)]},
        "idempotency_key": "tracker-replay-denial-001",
    }
    first = await f.call("set_travel_tracker_status", **args)
    assert "receipt" in first
    grant = await f.session.get(MCPGrantModel, f.principal.grant_id)
    grant.allowed_write_sections = ["all_groups"]
    await f.session.commit()
    assert (await f.call("set_travel_tracker_status", **args))["error"] == "access_denied"
    await f.session.refresh(f.user)
    f.token, f.principal = await f.connect(["mcp:read"])
    assert (
        await f.call(
            "set_travel_tracker_status", **{**args, "idempotency_key": "tracker-readonly-scope-001"}
        )
    )["error"] == "access_denied"


async def test_actual_native_workbook_preview_and_apply_only_unique_matches(tracker):
    f = tracker
    book, output = Workbook(), io.BytesIO()
    book.active.title = "Progress"
    book.active.append(["Name"])
    book.active.append([f.passengers[0].client_name])
    book.active.append([f.passengers[1].client_name])
    book.active.append(["Missing person"])
    book.save(output)
    book.close()
    content = output.getvalue()
    ticket, data = await upload_ticket(
        f, data=content, purpose="group_workbook", key="tracker-native-source-001"
    )
    response = await put(f, ticket, data)
    assert response.status_code == 200, response.text
    staged = await f.native.inspect(f.token, uuid.UUID(ticket["id"]))
    await f.session.commit()
    draft = {
        "group_id": str(f.group.id),
        "upload_id": staged["staged_source"]["upload_id"],
        "source_sha256": hashlib.sha256(data).hexdigest(),
        "sheet_name": "Progress",
        "track": "flight",
        "marked": True,
    }
    preview = await f.call("preview_travel_tracker_workbook", draft=draft)
    assert (preview["matched_count"], preview["ambiguous_count"], preview["unmatched_count"]) == (
        1,
        1,
        1,
    ), preview
    args = {
        "command": {"draft": draft, "expected_preview_hash": preview["preview_hash"]},
        "idempotency_key": "tracker-workbook-apply-001",
    }
    applied = await f.call("apply_travel_tracker_workbook", **args)
    assert applied["receipt"]["data"]["updated_count"] == 1, applied
    replay = await f.call("apply_travel_tracker_workbook", **args)
    assert replay["receipt"] == applied["receipt"]
    wrong_group = await f.call(
        "preview_travel_tracker_workbook", draft={**draft, "group_id": str(uuid.uuid4())}
    )
    assert wrong_group["error"] == "tracker_workbook_unavailable"


async def test_export_actual_xlsx_and_current_documents_permission_at_download(tracker):
    f = tracker
    request = TrackerExportRequest(group_id=f.group.id, track="visa", status="pending")
    service = MCPTravelTrackerExportService(f.session, f.settings, artifacts=f.service)
    observed = await service.inspect(f.principal, request)
    definition = tracker_export_operation(f.settings)
    receipt = await MCPOperationService(f.session, f.settings, [definition]).execute(
        access_token=f.token,
        operation_name=definition.policy.name,
        idempotency_key="tracker-export-actual-001",
        payload={
            "export": request.model_dump(mode="json"),
            "expected_revision": observed["expected_revision"],
        },
    )
    await f.session.commit()
    result = await service.generate(
        access_token=f.token, operation_id=uuid.UUID(receipt["operation_id"])
    )
    await f.session.commit()
    content = await downloaded(f, result)
    book = load_workbook(io.BytesIO(content), read_only=True)
    rows = list(book.active.values)
    assert "Passenger ID" in rows[0] and "Visa Applied" in rows[0] and "Flight Booked" in rows[0]
    assert len(rows) == 4
    book.close()
    recovered = await service.generate(
        access_token=f.token, operation_id=uuid.UUID(receipt["operation_id"])
    )
    await f.session.commit()
    assert recovered["artifact"]["sha256"] == result["artifact"]["sha256"]
    download = await f.native.create_download(
        f.token, result["artifact"]["artifact_id"], "tracker-download-001"
    )
    await f.session.commit()
    assert download["kind"] == "download"
    grant = await f.session.get(MCPGrantModel, f.principal.grant_id)
    grant.allowed_write_sections = ["exports"]
    await f.session.commit()
    with pytest.raises(MCPAuthError):
        await f.service.get(f.principal, result["artifact"]["artifact_id"])
    await f.session.rollback()


async def test_mcp_export_tool_roundtrip_and_revision_change(tracker):
    f = tracker
    export = {"group_id": str(f.group.id)}
    observed = await f.call("inspect_travel_tracker_export", export=export)
    assert observed["passenger_count"] == 3, observed
    result = await f.call(
        "prepare_travel_tracker_export",
        export=export,
        expected_revision=observed["expected_revision"],
        idempotency_key="tracker-sdk-export-001",
    )
    assert result["status"] == "succeeded" and "artifact" in result, result
    await f.call(
        "set_travel_tracker_status",
        group_id=str(f.group.id),
        update={"track": "visa", "marked": True, "passenger_ids": [str(f.passengers[0].id)]},
        idempotency_key="tracker-export-revision-change-001",
    )
    stale = await f.call(
        "prepare_travel_tracker_export",
        export=export,
        expected_revision=observed["expected_revision"],
        idempotency_key="tracker-sdk-stale-export-001",
    )
    assert stale["error"] == "tracker_export_changed"


async def test_release_registry_and_section_catalog_are_explicit(tracker):
    f = tracker
    assert READ_TOOL_SECTIONS["get_travel_tracker_roster"] == frozenset({"documents", "all_groups"})
    assert WRITE_TOOL_SECTIONS["prepare_travel_tracker_export"] == frozenset(
        {"documents", "all_groups", "exports"}
    )
    read_only = ObservationalMCPServer("Read only fixture")
    register_travel_tracker_tools(read_only, f.app, f.settings)
    assert {tool.name for tool in await read_only.list_tools()} == {
        "list_travel_tracker_groups",
        "get_travel_tracker_roster",
    }


@pytest.mark.parametrize("marked", ["true", "false", 1])
async def test_sdk_rejects_coerced_mark_boolean_without_effects(tracker, marked):
    f = tracker
    with pytest.raises(ToolError):
        await f.call(
            "set_travel_tracker_status",
            group_id=str(f.group.id),
            update={"track": "visa", "marked": marked, "passenger_ids": [str(f.passengers[0].id)]},
            idempotency_key="tracker-invalid-boolean-001",
        )
    assert await f.session.scalar(select(func.count()).select_from(TravelTrackerModel)) == 0


async def test_mcp_export_admits_rows_and_json_bytes_before_orm_values(tracker):
    f = tracker
    bounded = f.settings.model_copy(
        update={"mcp": f.settings.mcp.model_copy(update={"export_source_row_limit": 2})}
    )
    service = MCPTravelTrackerExportService(f.session, bounded, artifacts=f.service)
    with pytest.raises(ArtifactError, match="row limit"):
        await service.inspect(f.principal, TrackerExportRequest(group_id=f.group.id))
    await f.session.rollback()
    await f.session.refresh(f.group)
    await f.session.refresh(f.passengers[0])
    f.passengers[0].confirmed_fields = {"passport_number": "X" * 18000}
    await f.session.commit()
    bounded = f.settings.model_copy(
        update={"mcp": f.settings.mcp.model_copy(update={"export_source_byte_limit": 1024})}
    )
    service = MCPTravelTrackerExportService(f.session, bounded, artifacts=f.service)
    with pytest.raises(ArtifactError, match="byte limit"):
        await service.inspect(f.principal, TrackerExportRequest(group_id=f.group.id))
