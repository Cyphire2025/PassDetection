"""Read-only superset of the normal passport export for broadcast filter rows."""

from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from typing import Any

from openpyxl import Workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE, Cell
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from app.domain.entities.entities import PassportSubmission
from app.domain.value_objects.client_collection_provenance import CLIENT_COLLECTION_SUBMITTED_KEY
from app.infrastructure.export.passport_excel_exporter import (
    PassportExcelExporter,
    _excel_date_value,
)

MAX_EXPORT_ROWS = 20_000
MAX_EXPORT_COLUMNS = 1_024
MAX_EXPORT_CELLS = 2_000_000
MAX_EXPORT_TEXT_CHARACTERS = 32_000_000
MAX_EXPORT_TEXT_BYTES = 64 * 1024 * 1024
_DATE_COLUMNS = {"DOB", "DOI", "DOE", "Travel/Departure Date", "Return Date"}
_XML_INVALID_UNICODE = re.compile(r"[\ud800-\udfff\ufffe\uffff]")
_PRIMARY_COLUMNS = (
    "Broadcast contact name",
    "Broadcast phone",
    "Broadcast normalized phone",
    "Source contact name",
    "Source phone",
    "Source normalized phone",
    "Source issue",
    "Roster row type",
    "Filter",
    "Filter message type",
    "Selected message status",
)
_CUSTOM_FIELD_SOURCES = (
    ("Custom question", "custom_questions", "custom_answers", "question_id"),
    ("Custom detail", "custom_details", "custom_detail_answers", "detail_id"),
)


class WhatsAppExportTooLarge(ValueError):
    """The requested complete export cannot fit its bounded workbook budget."""


@dataclass(frozen=True)
class WhatsAppExportRow:
    values: dict[str, Any]
    broadcast_fields: dict[str, Any] = field(default_factory=dict)
    submission: PassportSubmission | None = None
    group_details: dict[str, Any] = field(default_factory=dict)


def _custom_columns(rows: list[WhatsAppExportRow]) -> dict[tuple[str, str], str]:
    """Keep one column per stable identity, including unanswered configured fields."""
    columns: dict[tuple[str, str], str] = {}
    seen_groups: set[str] = set()

    def include(source: str, identity: str, label: str) -> None:
        columns.setdefault((source, identity), f"{source}: {label} [{identity}]")
        if len(columns) + len(PassportExcelExporter.HEADERS) > MAX_EXPORT_COLUMNS:
            raise WhatsAppExportTooLarge(
                "This export contains too many saved fields. "
                "Narrow the filter or export one source group at a time."
            )

    for row in rows:
        if row.submission is None or str(row.submission.group_id) in seen_groups:
            continue
        seen_groups.add(str(row.submission.group_id))
        for source, definitions_key, _, _ in _CUSTOM_FIELD_SOURCES:
            for definition in row.group_details.get(definitions_key) or []:
                if not definition.get("enabled") or not definition.get("id"):
                    continue
                identity = str(definition["id"])
                label = str(definition.get("label") or identity)
                include(source, identity, label)
    # Saved answers to disabled/removed fields remain exportable. Current
    # definitions take precedence over historical labels for the same identity.
    for row in rows:
        if row.submission is None:
            continue
        for source, _, answers_key, identity_key in _CUSTOM_FIELD_SOURCES:
            for index, answer in enumerate(getattr(row.submission, answers_key) or [], start=1):
                identity = str(answer.get(identity_key) or index)
                label = str(answer.get("label") or identity)
                include(source, identity, label)
    return columns


def _supplemental_values(
    row: WhatsAppExportRow,
    custom_columns: dict[tuple[str, str], str],
) -> dict[str, Any]:
    values: dict[str, Any] = {}
    submission = row.submission
    if submission is not None:
        # Reuse the normal export's authoritative map and human-readable values,
        # but include every field irrespective of the source group's UI toggles.
        values.update(PassportExcelExporter._submission_values(
            submission, group_name=str(row.group_details.get("name") or ""),
            details=row.group_details, zone_names=None, dynamic_fields=[],
            additional_values=None, whatsapp_contacts=None, previous_names=None,
        ))
        values.update({
            "Source submission ID": str(submission.id),
            "Source group ID": str(submission.group_id),
            "Saved client name": submission.client_name,
            "Saved client email": submission.client_email,
            "Saved client phone": submission.client_phone,
            "Submission status": submission.status.value,
            "Submission mode": submission.submission_mode,
            "Family head name": submission.family_head_name,
            "Family head email": submission.family_head_email,
            "Family head phone": submission.family_head_phone,
            "Family relation": submission.family_relation,
            "Family gender": submission.family_gender,
            "Family member order": submission.family_member_index,
            "Family broadcast to member": submission.family_broadcast_to_member,
            "Source created at": submission.created_at,
            "Source updated at": submission.updated_at,
            "Package": row.group_details.get("package_name"),
        })
        # Do not overlay missing/cleared confirmed values with stale OCR fields.
        for key, value in (submission.confirmed_fields or submission.extracted_fields or {}).items():
            values[f"Saved passport: {key}"] = value
        for key, value in (submission.staff_metadata or {}).items():
            if key != CLIENT_COLLECTION_SUBMITTED_KEY:
                values[f"Source import: {key}"] = value
        for source, _, answers_key, identity_key in _CUSTOM_FIELD_SOURCES:
            for index, answer in enumerate(getattr(submission, answers_key) or [], start=1):
                identity = str(answer.get(identity_key) or index)
                values[custom_columns[(source, identity)]] = answer.get("value")
    for key, value in row.broadcast_fields.items():
        values[f"Broadcast import: {key}"] = value
    values.update(row.values)
    return values


def _text_value(value: Any) -> str:
    if isinstance(value, datetime):
        text = value.isoformat()
    elif isinstance(value, (dict, list, tuple)):
        text = json.dumps(value, ensure_ascii=False, default=str)
    else:
        text = str(value)
    # Escape XML-incompatible controls visibly instead of losing saved data.
    text = ILLEGAL_CHARACTERS_RE.sub(lambda match: f"\\x{ord(match.group()):02x}", text)
    text = _XML_INVALID_UNICODE.sub(lambda match: f"\\u{ord(match.group()):04x}", text)
    if len(text) > 32_767:
        raise WhatsAppExportTooLarge(
            "A saved field exceeds Excel's 32,767-character cell limit. "
            "Narrow the filter or shorten that source field before exporting."
        )
    return text


@dataclass
class _ExportTextBudget:
    characters: int = 0
    utf8_bytes: int = 0

    def include(self, text: str) -> None:
        self.characters += len(text)
        self.utf8_bytes += len(text.encode("utf-8"))
        if (
            self.characters > MAX_EXPORT_TEXT_CHARACTERS
            or self.utf8_bytes > MAX_EXPORT_TEXT_BYTES
        ):
            raise WhatsAppExportTooLarge(
                "This export contains too much saved text. "
                "Narrow the filter or export one source group at a time."
            )


def build_whatsapp_filter_workbook(rows: list[WhatsAppExportRow]) -> bytes:
    """Only detached values/entities enter this function; it runs in a worker thread."""
    if len(rows) > MAX_EXPORT_ROWS:
        raise WhatsAppExportTooLarge(
            "This filter expands to more than 20,000 traveller rows. Narrow the filter and export again."
        )
    custom_columns = _custom_columns(rows)
    headers = list(dict.fromkeys([*PassportExcelExporter.HEADERS, *custom_columns.values()]))
    known_headers = set(headers)
    header_text: dict[str, str] = {}
    budget = _ExportTextBudget()
    prepared: list[dict[str, str | date]] = []

    def include_header(header: str) -> None:
        text = _text_value(header)
        budget.include(text)
        header_text[header] = text

    for header in headers:
        include_header(header)
    for row in rows:
        values = _supplemental_values(row, custom_columns)
        for header in values:
            if header not in known_headers:
                known_headers.add(header)
                headers.append(header)
                include_header(header)
        if len(headers) > MAX_EXPORT_COLUMNS or (len(rows) + 1) * len(headers) > MAX_EXPORT_CELLS:
            raise WhatsAppExportTooLarge(
                "This export contains too many saved fields or cells. "
                "Narrow the filter or export one source group at a time."
            )
        prepared_values: dict[str, str | date] = {}
        for header, value in values.items():
            if value is None:
                continue
            if header in _DATE_COLUMNS:
                value = _excel_date_value(value)
            if isinstance(value, date) and not isinstance(value, datetime):
                budget.include(value.isoformat())
                prepared_values[header] = value
            else:
                text = _text_value(value)
                budget.include(text)
                prepared_values[header] = text
        prepared.append(prepared_values)
    headers = [header for header in _PRIMARY_COLUMNS if header in known_headers] + [
        header for header in headers if header not in _PRIMARY_COLUMNS
    ]

    # Validate and serialize values before opening the streaming XML writer so
    # size errors never leave a partial workbook or silently truncated cells.
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Broadcast export")
    sheet.sheet_view.showGridLines = True
    sheet.freeze_panes = "A2"
    fill = PatternFill("solid", fgColor="165D52")
    header_font = Font(bold=True, color="FFFFFF")
    header_alignment = Alignment(vertical="center", wrap_text=True)
    body_alignment = Alignment(vertical="top", wrap_text=True)
    sheet.row_dimensions[1].height = 32
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{len(prepared) + 1}"
    header_cells: list[Cell] = []
    for column, header in enumerate(headers, start=1):
        cell = WriteOnlyCell(sheet, value=header_text[header])
        cell.data_type = "s"
        cell.font = header_font
        cell.fill = fill
        cell.alignment = header_alignment
        sheet.column_dimensions[get_column_letter(column)].width = min(42, max(18, len(header) + 2))
        header_cells.append(cell)
    sheet.append(header_cells)
    for values in prepared:
        cells: list[Cell | None] = []
        for header in headers:
            value = values.get(header)
            if value is None:
                cells.append(None)
                continue
            cell = WriteOnlyCell(sheet, value=value)
            if isinstance(value, date):
                cell.number_format = "DD.MM.YYYY"
            else:
                # Explicit strings prevent formulas and preserve +country codes,
                # leading zeros and long identifiers without changing their value.
                cell.data_type = "s"
                cell.number_format = "@"
            cell.alignment = body_alignment
            cells.append(cell)
        sheet.append(cells)
    output = io.BytesIO()
    try:
        workbook.save(output)
    finally:
        workbook.close()
    return output.getvalue()
