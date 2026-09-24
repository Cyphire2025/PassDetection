"""Two-column ECR export with literal filenames and explicit unresolved rows."""

from collections.abc import Sequence
from io import BytesIO

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Font, PatternFill


def ecr_result_font(result: str) -> Font:
    """Keep standalone and passport-group ECR verdict colors identical."""
    return Font(
        color="FF0000" if result == "ECR" else "000000" if result == "NA" else "B45309",
        bold=result == "ECR",
    )


def build_ecr_workbook(rows: Sequence[tuple[str, str]]) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "ECR"
    sheet.append(["File name", "ECR"])
    for cell in sheet[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="164E63")
    for filename, result in rows:
        if result not in {"ECR", "NA", "REVIEW", "ERROR", "PENDING"}:
            result = "REVIEW"
        sheet.append([ILLEGAL_CHARACTERS_RE.sub("", filename), result])
        # String cells preserve exact names, including '=' prefixes, without formulas.
        sheet.cell(sheet.max_row, 1).data_type = "s"
        sheet.cell(sheet.max_row, 2).font = ecr_result_font(result)
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    sheet.column_dimensions["A"].width = 64
    sheet.column_dimensions["B"].width = 18
    output = BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()
