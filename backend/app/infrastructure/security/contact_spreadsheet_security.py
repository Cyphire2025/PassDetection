"""Fixed XLSX ingestion using the existing malware/quarantine evidence boundary."""

from __future__ import annotations

import asyncio
import json
from io import BytesIO
from itertools import islice
from typing import Any
from xml.etree import ElementTree as ET
from zipfile import ZipFile

from openpyxl import load_workbook

from app.application.use_cases.whatsapp.spreadsheet_values import (
    SpreadsheetArchiveCapacityError,
    excel_cell_text,
    validate_excel_archive,
)
from app.domain.contact_workbook import ContactWorkbookValidationError
from app.domain.exceptions.exceptions import ImageValidationError
from app.infrastructure.documents.storage_transfers import run_bounded_storage_operations
from app.infrastructure.security.upload_security import UploadSecurityContext, UploadSecurityService

XLSX_MEDIA = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
MAX_BYTES = 5 * 1024 * 1024
MAX_SNAPSHOT_BYTES = 2 * 1024 * 1024
MAX_ROWS, MAX_COLUMNS, MAX_CELL_CHARS, MAX_CELLS = 2000, 64, 1024, 50000


class _NoDTD(ET.TreeBuilder):
    def doctype(self, name: str, pubid: str | None, system: str | None) -> None:
        raise ContactWorkbookValidationError("contact_workbook_active_content_unsupported")


def _safe_archive(content: bytes) -> None:
    try:
        validate_excel_archive(content)
    except SpreadsheetArchiveCapacityError:
        raise ContactWorkbookValidationError("contact_workbook_capacity_exceeded") from None
    with ZipFile(BytesIO(content)) as archive:
        names = archive.namelist()
        if len(names) != len(set(names)) or "[Content_Types].xml" not in names:
            raise ContactWorkbookValidationError("contact_workbook_invalid_package")
        for member in archive.infolist():
            name = member.filename.lower()
            if member.flag_bits & 1 or any(
                part in name for part in ("vbaproject", "externallink", "activex", "embeddings/")
            ):
                raise ContactWorkbookValidationError("contact_workbook_active_content_unsupported")
            if not name.endswith((".xml", ".rels")):
                continue
            root = ET.fromstring(archive.read(member), parser=ET.XMLParser(target=_NoDTD()))
            for element in root.iter():
                tag = element.tag.rsplit("}", 1)[-1]
                if tag in {"f", "definedName"} or "formula" in tag.casefold():
                    raise ContactWorkbookValidationError("contact_workbook_formulas_unsupported")
                if element.get("TargetMode", "").casefold() == "external":
                    raise ContactWorkbookValidationError(
                        "contact_workbook_active_content_unsupported"
                    )
                if "macroenabled" in element.get("ContentType", "").casefold():
                    raise ContactWorkbookValidationError(
                        "contact_workbook_active_content_unsupported"
                    )


def snapshot_workbook(content: bytes) -> dict[str, Any]:
    """Keep bounded literal cells without guessing headers or silently dropping values."""
    _safe_archive(content)
    workbook = load_workbook(BytesIO(content), read_only=True, data_only=False, keep_links=False)
    try:
        if not 1 <= len(workbook.worksheets) <= 50:
            raise ContactWorkbookValidationError("contact_workbook_capacity_exceeded")
        sheets, total_rows, total_cells, text_bytes = [], 0, 0, 0
        for sheet in workbook.worksheets:
            # Do not trust dimension metadata supplied in the ZIP package.
            sheet.reset_dimensions()
            rows = []
            for row in islice(sheet.iter_rows(), MAX_ROWS - total_rows + 1):
                total_rows += 1
                total_cells += len(row)
                if total_rows > MAX_ROWS or len(row) > MAX_COLUMNS or total_cells > MAX_CELLS:
                    raise ContactWorkbookValidationError("contact_workbook_capacity_exceeded")
                values = []
                for cell in row:
                    if cell.data_type == "f":
                        raise ContactWorkbookValidationError(
                            "contact_workbook_formulas_unsupported"
                        )
                    value = excel_cell_text(cell.value)
                    text_bytes += len(value.encode("utf-8"))
                    if len(value) > MAX_CELL_CHARS or text_bytes > MAX_SNAPSHOT_BYTES:
                        raise ContactWorkbookValidationError("contact_workbook_capacity_exceeded")
                    values.append(value)
                rows.append(values)
            sheets.append({"name": sheet.title, "rows": rows})
        result = {"schema_version": 1, "sheets": sheets, "row_count": total_rows}
        if len(json.dumps(result, ensure_ascii=False).encode("utf-8")) > MAX_SNAPSHOT_BYTES:
            raise ContactWorkbookValidationError("contact_workbook_capacity_exceeded")
        return result
    finally:
        workbook.close()


class ContactSpreadsheetSecurity(UploadSecurityService):
    async def validate_spreadsheet(
        self, *, content: bytes, context: UploadSecurityContext
    ) -> dict[str, Any]:
        if not 1 <= len(content) <= MAX_BYTES:
            raise ImageValidationError("Contact workbook size is invalid")
        await run_bounded_storage_operations(
            [
                lambda: self._scan_original(
                    content=content, declared_media_type=XLSX_MEDIA, context=context
                )
            ],
            concurrency=1,
        )
        try:
            # Retain the transfer slot until the parser exits after cancellation.
            snapshot = (
                await run_bounded_storage_operations(
                    [lambda: asyncio.to_thread(snapshot_workbook, content)], concurrency=1
                )
            )[0]
        except Exception as exc:
            await self._record(
                content=content,
                declared_media_type=XLSX_MEDIA,
                context=context,
                scan_status="malformed",
                disposition="rejected",
                error_code="XLSX_VALIDATION_FAILED",
            )
            if isinstance(exc, ContactWorkbookValidationError):
                raise ContactWorkbookValidationError(exc.code) from None
            raise ContactWorkbookValidationError("contact_workbook_invalid_package") from None
        await self._record(
            content=content,
            declared_media_type=XLSX_MEDIA,
            context=context,
            scan_status="clean",
            disposition="accepted",
        )
        return snapshot
