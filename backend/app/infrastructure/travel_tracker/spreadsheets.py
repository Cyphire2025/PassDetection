"""Bounded spreadsheets with conservative, reviewable identity matching."""

from __future__ import annotations

import csv
import io
import json
import re
import unicodedata
import uuid
from collections import defaultdict
from collections.abc import Iterable
from typing import Any
from zipfile import ZipFile

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

from app.application.dtos.travel_tracker import (
    TrackerImportPreview,
    TrackerImportRow,
    TrackerTrack,
)
from app.infrastructure.database.models import ClientGroupModel, PassportSubmissionModel
from app.infrastructure.database.travel_tracker_model import TravelTrackerModel

MAX_UPLOAD_BYTES = 8 * 1024 * 1024
MAX_IMPORT_ROWS = 1000
MAX_EXPORT_ROWS = 20000
MAX_COLUMNS = 256


class TrackerSpreadsheetError(ValueError):
    pass


def field_values(passenger: PassportSubmissionModel) -> dict[str, Any]:
    return {**(passenger.extracted_fields or {}), **(passenger.confirmed_fields or {})}


def text_value(value: Any) -> str | None:
    if value is None:
        return None
    if isinstance(value, (dict, list)):
        return json.dumps(value, ensure_ascii=False, sort_keys=True)
    return str(value).strip() or None


def full_name(passenger: PassportSubmissionModel) -> str:
    fields = field_values(passenger)
    return (
        " ".join(str(passenger.client_name or "").split())
        or " ".join(
            str(fields.get(key) or "").strip() for key in ("given_names", "surname")
        ).strip()
        or "Unnamed passenger"
    )


def normalized_name(value: Any) -> str:
    return " ".join(unicodedata.normalize("NFKC", str(value or "")).split()).casefold()


def normalized_passport(value: Any) -> str:
    return re.sub(r"[\s-]+", "", unicodedata.normalize("NFKC", str(value or ""))).upper()


def _header(value: Any) -> str:
    return re.sub(r"[^a-z0-9]", "", normalized_name(value))


_HEADERS = {
    "passengerid": "id",
    "submissionid": "id",
    "id": "id",
    "passportnumber": "passport",
    "passportno": "passport",
    "passport": "passport",
    "name": "name",
    "fullname": "name",
    "passengername": "name",
    "clientname": "name",
    "travellername": "name",
    "travelername": "name",
    "givenname": "given",
    "givennames": "given",
    "firstname": "given",
    "surname": "surname",
    "lastname": "surname",
    "familyname": "surname",
}


def _bounded_rows(rows: Iterable[Iterable[Any]]) -> list[dict[str, Any]]:
    columns: dict[str, int] | None = None
    output: list[dict[str, Any]] = []
    for row_number, raw in enumerate(rows, 1):
        if row_number > MAX_IMPORT_ROWS + 20:
            raise TrackerSpreadsheetError("Import at most 1,000 passengers at a time.")
        values = list(raw)
        if len(values) > MAX_COLUMNS:
            raise TrackerSpreadsheetError("Spreadsheets can contain at most 256 columns.")
        if not any(value not in (None, "") for value in values):
            continue
        if any(len(str(value or "")) > 2048 for value in values):
            raise TrackerSpreadsheetError("Spreadsheet cells can contain at most 2,048 characters.")
        if columns is None:
            candidates = {
                _HEADERS[_header(value)]: index
                for index, value in enumerate(values)
                if _header(value) in _HEADERS
            }
            if (
                "id" in candidates
                or "passport" in candidates
                or "name" in candidates
                or "given" in candidates
            ):
                columns = candidates
                continue
            if row_number >= 20:
                break
            continue
        if len(output) >= MAX_IMPORT_ROWS:
            raise TrackerSpreadsheetError("Import at most 1,000 passengers at a time.")
        record = {
            key: values[index] if index < len(values) else None for key, index in columns.items()
        }
        record["row_number"] = row_number
        if not record.get("name"):
            record["name"] = " ".join(
                str(record.get(key) or "").strip() for key in ("given", "surname")
            ).strip()
        output.append(record)
    if columns is None:
        raise TrackerSpreadsheetError(
            "Add a Name, Passenger ID, or Passport Number column to the file."
        )
    if not output:
        raise TrackerSpreadsheetError("The spreadsheet has no passenger rows.")
    return output


def read_rows(content: bytes, filename: str) -> list[dict[str, Any]]:
    if not content:
        raise TrackerSpreadsheetError("The uploaded file is empty.")
    if len(content) > MAX_UPLOAD_BYTES:
        raise TrackerSpreadsheetError("Upload a file smaller than 8 MB.")
    if filename.lower().endswith(".csv"):
        try:
            source = content.decode("utf-8-sig")
            return _bounded_rows(csv.reader(io.StringIO(source)))
        except (UnicodeDecodeError, csv.Error) as exc:
            raise TrackerSpreadsheetError("Save the CSV using UTF-8 and try again.") from exc
    if not filename.lower().endswith(".xlsx"):
        raise TrackerSpreadsheetError("Upload an .xlsx workbook or a UTF-8 .csv file.")
    try:
        with ZipFile(io.BytesIO(content)) as archive:
            entries = archive.infolist()
            if len(entries) > 2000 or sum(entry.file_size for entry in entries) > 32 * 1024 * 1024:
                raise TrackerSpreadsheetError("The workbook expands beyond the allowed size.")
            for entry in entries:
                path = entry.filename.lower()
                if "vbaproject" in path or "externallinks/" in path or entry.flag_bits & 1:
                    raise TrackerSpreadsheetError(
                        "Remove macros, external workbook links, and encryption before uploading."
                    )
                if (
                    entry.file_size > 1024 * 1024
                    and entry.file_size / max(entry.compress_size, 1) > 250
                ):
                    raise TrackerSpreadsheetError("The workbook has an unsafe compression ratio.")
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=False, keep_links=False
        )
        try:
            if len(workbook.worksheets) != 1:
                raise TrackerSpreadsheetError(
                    "Use a workbook with one sheet so the update list is unambiguous."
                )
            sheet = workbook.worksheets[0]
            if (sheet.max_column or 0) > MAX_COLUMNS:
                raise TrackerSpreadsheetError("Spreadsheets can contain at most 256 columns.")
            # Ignore declared dimensions and stream actual rows under our limit.
            sheet.reset_dimensions()
            return _bounded_rows(sheet.iter_rows(values_only=True))
        finally:
            workbook.close()
    except TrackerSpreadsheetError:
        raise
    except Exception as exc:
        # Invalid OOXML raises different parser exceptions depending on the
        # optional XML backend. This bounded parser boundary must fail closed.
        raise TrackerSpreadsheetError(
            "The workbook could not be read. Save it as .xlsx and try again."
        ) from exc


def match_rows(
    records: list[dict[str, Any]],
    passengers: list[PassportSubmissionModel],
    *,
    track: TrackerTrack,
    marked: bool,
) -> TrackerImportPreview:
    by_id = {str(passenger.id): passenger for passenger in passengers}
    by_passport: dict[str, dict[uuid.UUID, PassportSubmissionModel]] = defaultdict(dict)
    by_name: dict[str, dict[uuid.UUID, PassportSubmissionModel]] = defaultdict(dict)
    for passenger in passengers:
        fields = field_values(passenger)
        normalized_number = normalized_passport(fields.get("passport_number"))
        if normalized_number:
            by_passport[normalized_number][passenger.id] = passenger
        for alias_name in (
            full_name(passenger),
            " ".join(str(fields.get(key) or "").strip() for key in ("given_names", "surname")),
        ):
            if normalized_name(alias_name):
                by_name[normalized_name(alias_name)][passenger.id] = passenger
    result: list[TrackerImportRow] = []
    matched_ids: list[uuid.UUID] = []
    seen: set[uuid.UUID] = set()
    for record in records:
        supplied_id = str(record.get("id") or "").strip()
        name = text_value(record.get("name"))
        passport = text_value(record.get("passport"))
        row = TrackerImportRow(
            row_number=record["row_number"],
            name=name,
            passport_number=passport,
            status="unmatched",
            reason="No passenger has this exact name in this group.",
        )
        candidates: list[PassportSubmissionModel] = []
        if any(
            str(record.get(key) or "").lstrip().startswith("=")
            for key in ("id", "name", "passport")
        ):
            row.reason = "Identity cells must contain text, not spreadsheet formulas."
        elif supplied_id:
            try:
                resolved_id = str(uuid.UUID(supplied_id))
            except ValueError:
                resolved_id = ""
            candidate = by_id.get(resolved_id)
            if candidate is not None and (
                not passport
                or normalized_passport(passport)
                == normalized_passport(field_values(candidate).get("passport_number"))
            ):
                candidates = [candidate]
                row.reason = "Matched Passenger ID."
            else:
                row.reason = (
                    "Passenger ID is missing from this group or conflicts with the passport number."
                )
        elif passport:
            candidates = list(by_passport.get(normalized_passport(passport), {}).values())
            row.reason = (
                "Matched passport number."
                if candidates
                else "Passport number was not found in this group."
            )
        elif name:
            candidates = list(by_name.get(normalized_name(name), {}).values())
            if candidates:
                row.reason = "Matched exact unique name."
        if len(candidates) > 1:
            row.status = "ambiguous"
            row.reason = "Multiple passengers match. Add a Passenger ID or unique Passport Number."
        elif len(candidates) == 1:
            passenger = candidates[0]
            row.passenger_id, row.passenger_name = passenger.id, full_name(passenger)
            if passenger.id in seen:
                row.status, row.reason = (
                    "duplicate",
                    "This passenger already appears in the update list.",
                )
            else:
                row.status = "matched"
                seen.add(passenger.id)
                matched_ids.append(passenger.id)
        result.append(row)
    counts = {
        kind: sum(row.status == kind for row in result)
        for kind in ("matched", "ambiguous", "unmatched", "duplicate")
    }
    return TrackerImportPreview(
        track=track,
        marked=marked,
        total_rows=len(result),
        matched_count=counts["matched"],
        ambiguous_count=counts["ambiguous"],
        unmatched_count=counts["unmatched"],
        duplicate_count=counts["duplicate"],
        passenger_ids=matched_ids,
        rows=result,
    )


def _safe(value: Any) -> Any:
    if isinstance(value, (dict, list)):
        value = json.dumps(value, ensure_ascii=False, sort_keys=True)
    if isinstance(value, str):
        value = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f]", "", value)[:32767]
        if value.lstrip().startswith(("=", "+", "-", "@")):
            return "'" + value
    return value


def build_export(
    group: ClientGroupModel, rows: list[tuple[PassportSubmissionModel, TravelTrackerModel | None]]
) -> bytes:
    """Include fixed contact/trip details and every collected/imported data field."""
    headers = [
        "Passenger ID",
        "Full Name",
        "Group",
        "Destination",
        "Travel Date",
        "Return Date",
        "Package",
        "Email",
        "Phone",
        "Departure City",
        "Nearest Domestic Airport",
        "Submission Status",
        "Visa Applied",
        "Flight Booked",
        "Visa Updated At",
        "Flight Updated At",
        "Visa Updated By",
        "Flight Updated By",
        "Family Relation",
        "Family Head Name",
        "Relation With Qualifier",
    ]
    passport_keys = sorted({key for passenger, _ in rows for key in field_values(passenger)})
    # Passport Number remains recognizable by round-trip identity matching.
    passport_headers = {
        key: " ".join(word.title() for word in key.split("_")) for key in passport_keys
    }
    metadata_keys = sorted(
        {
            key
            for passenger, _ in rows
            for key in (passenger.staff_metadata or {})
            if not key.startswith("_")
        }
    )
    custom_questions = {
        str(item.get("id")): str(item.get("label") or item.get("question") or item.get("id"))
        for item in (group.custom_questions or [])
    }
    custom_details = {
        str(item.get("id")): str(item.get("label") or item.get("id"))
        for item in (group.custom_details or [])
    }
    for passenger, _ in rows:
        for answer in passenger.custom_answers or []:
            custom_questions.setdefault(
                str(answer.get("question_id")),
                str(answer.get("label") or answer.get("question_id")),
            )
        for detail_answer in passenger.custom_detail_answers or []:
            custom_details.setdefault(
                str(detail_answer.get("detail_id")),
                str(detail_answer.get("label") or detail_answer.get("detail_id")),
            )
    headers += [passport_headers[key] for key in passport_keys]
    headers += ["Imported: " + key for key in metadata_keys]
    headers += ["Question: " + label for label in custom_questions.values()]
    headers += ["Detail: " + label for label in custom_details.values()]
    if len(headers) > MAX_COLUMNS:
        raise TrackerSpreadsheetError(
            "The roster has too many detail columns for one export (maximum 256)."
        )
    workbook = Workbook(write_only=True)
    sheet = workbook.create_sheet("Travel Tracker")
    sheet.freeze_panes = "C2"
    for index, header in enumerate(headers, 1):
        sheet.column_dimensions[get_column_letter(index)].width = min(38, max(18, len(header) + 3))

    styled_headers = []
    for value in headers:
        cell = WriteOnlyCell(sheet, value=_safe(value))
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="153A56")
        styled_headers.append(cell)
    sheet.append(styled_headers)
    for passenger, tracker in rows:
        fields = field_values(passenger)
        answers = {
            str(answer.get("question_id")): answer.get("value")
            for answer in passenger.custom_answers or []
        }
        details = {
            str(answer.get("detail_id")): answer.get("value")
            for answer in passenger.custom_detail_answers or []
        }
        values = [
            str(passenger.id),
            full_name(passenger),
            group.name,
            group.destination,
            group.travel_date,
            group.return_date,
            group.package_name,
            passenger.client_email,
            passenger.client_phone,
            passenger.departure_city,
            passenger.nearest_domestic_airport,
            passenger.status,
            bool(tracker and tracker.visa_applied),
            bool(tracker and tracker.flight_booked),
            tracker.visa_updated_at.isoformat() if tracker and tracker.visa_updated_at else None,
            tracker.flight_updated_at.isoformat()
            if tracker and tracker.flight_updated_at
            else None,
            str(tracker.visa_updated_by) if tracker and tracker.visa_updated_by else None,
            str(tracker.flight_updated_by) if tracker and tracker.flight_updated_by else None,
            passenger.family_relation,
            passenger.family_head_name,
            passenger.qualifier_relation_label,
        ]
        values += [fields.get(key) for key in passport_keys]
        values += [(passenger.staff_metadata or {}).get(key) for key in metadata_keys]
        values += [answers.get(key) for key in custom_questions]
        values += [details.get(key) for key in custom_details]
        sheet.append([_safe(value) for value in values])
    sheet.auto_filter.ref = f"A1:{get_column_letter(len(headers))}{max(1, len(rows) + 1)}"
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()
